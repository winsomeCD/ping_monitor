#!/usr/bin/env python3
"""
Ping Check Enhanced - Multi-Host Network Monitor (Pure PySide6)
Professional Desktop Edition
"""

import csv
import json
import math
import os
import re
import sys
import tempfile
import time
from collections import deque
from pathlib import Path

from PySide6.QtCore import Qt, Slot, QTimer, QPointF, QRectF, QThreadPool
from PySide6.QtGui import QBrush, QColor, QFont, QPainter, QPen
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QLabel, QVBoxLayout, QHBoxLayout,
    QGridLayout, QSplitter, QFileDialog, QAbstractItemView, QHeaderView,
    QTableWidget, QTableWidgetItem, QLineEdit, QDialog, QTextEdit,
    QPushButton, QSpinBox, QDoubleSpinBox, QFrame, QMessageBox,
)
from probes import MAX_CONCURRENT_PROBES, ProbeResult, ProbeTask, normalize_targets


_NUMERIC_RE = re.compile(r"[^0-9.+-]")


def _sanitize_csv_cell(value: object) -> object:
    """Neutralize spreadsheet formula injection (DDE attacks) for exported CSVs."""
    if value is None:
        return ""
    text = str(value)
    if text and text[0] in ("=", "+", "-", "@", "\t", "\r"):
        try:
            float(text)
            return value
        except ValueError:
            return "'" + text
    return value


class SortableTableWidgetItem(QTableWidgetItem):
    """Sort numbers by value using direct numeric data while preserving text sorting."""

    def __init__(self, text: str, numeric: bool = False, sort_value: float | None = None):
        super().__init__(text)
        self.numeric = numeric
        if sort_value is not None:
            self.setData(Qt.UserRole, sort_value)
        elif numeric:
            clean = _NUMERIC_RE.sub("", text)
            try:
                self.setData(Qt.UserRole, float(clean))
            except (ValueError, TypeError):
                self.setData(Qt.UserRole, float("inf") if text in ("—", "-", "n/a") else 0.0)

    def set_value(self, text: str, sort_value: float | None = None):
        if self.text() != text:
            self.setText(text)
        if sort_value is not None:
            self.setData(Qt.UserRole, sort_value)
        elif self.numeric:
            clean = _NUMERIC_RE.sub("", text)
            try:
                self.setData(Qt.UserRole, float(clean))
            except (ValueError, TypeError):
                self.setData(Qt.UserRole, float("inf") if text in ("—", "-", "n/a") else 0.0)

    def __lt__(self, other: object) -> bool:
        if self.numeric and getattr(other, "numeric", False):
            v_self = self.data(Qt.UserRole)
            v_other = getattr(other, "data", lambda r: None)(Qt.UserRole)
            if v_self is not None and v_other is not None:
                try:
                    return float(v_self) < float(v_other)
                except (ValueError, TypeError):
                    pass
            # Fallback if UserRole not set
            s_text = self.text()
            o_text = getattr(other, "text", lambda: "")()
            left = _NUMERIC_RE.sub("", s_text)
            right = _NUMERIC_RE.sub("", o_text)
            try:
                return float(left) < float(right)
            except (ValueError, TypeError):
                pass
        return self.text().casefold() < str(getattr(other, "text", lambda: "")()).casefold()


# ---------------------------------------------------------------------------
# Host Detail Dialog
# ---------------------------------------------------------------------------
class HostDetailDialog(QDialog):
    def __init__(self, host: str, data: dict, dark: bool = True, parent: QWidget | None = None):
        super().__init__(parent)
        self.host = host
        self.data = data
        self._dark = dark
        self.setWindowTitle(f"Host Telemetry — {host}")
        self.resize(660, 560)
        self._build_ui()
        self._apply_style()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(16)

        header = QLabel(f"Target: {self.host}")
        header.setObjectName("detailTitle")
        layout.addWidget(header)

        stats_layout = QGridLayout()
        stats_layout.setSpacing(12)
        sent = self.data.get("sent", 0)
        received = self.data.get("received", 0)
        loss_pct = round(((sent - received) / sent) * 100, 1) if sent > 0 else None
        valid_latencies = [l for l in self.data.get("latencies", []) if l > 0]
        stats = [
            ("Packets Sent", sent),
            ("Packets Received", received),
            ("Packet Loss", f"{loss_pct}%" if loss_pct is not None else "n/a"),
            ("Recent Samples", len(valid_latencies)),
        ]
        if valid_latencies:
            stats.extend([
                ("Min Latency", f"{min(valid_latencies):.2f} ms"),
                ("Max Latency", f"{max(valid_latencies):.2f} ms"),
                ("Avg Latency", f"{sum(valid_latencies) / len(valid_latencies):.2f} ms"),
            ])
        for i, (label, value) in enumerate(stats):
            row, col = i // 2, (i % 2) * 2
            lbl = QLabel(f"{label}:")
            lbl.setObjectName("mutedLabel")
            val = QLabel(str(value))
            val.setObjectName("strongLabel")
            stats_layout.addWidget(lbl, row, col)
            stats_layout.addWidget(val, row, col + 1)
        layout.addLayout(stats_layout)

        status = self.data.get("last_status", "Unknown")
        dark = self._dark
        if status == "Connected":
            status_color = "#10b981" if dark else "#059669"
        elif status == "No ICMP Reply":
            status_color = "#f59e0b" if dark else "#d97706"
        else:
            status_color = "#f43f5e" if dark else "#e11d48"

        status_row = QHBoxLayout()
        status_label = QLabel("Current Status:")
        status_label.setObjectName("mutedLabel")
        status_value = QLabel(status)
        status_value.setStyleSheet(f"color: {status_color}; font-weight: 700;")
        status_row.addWidget(status_label)
        status_row.addWidget(status_value)
        status_row.addStretch()
        layout.addLayout(status_row)

        if self.data.get("last_down_time") and self.data.get("last_down_time") != "-":
            down_lbl = QLabel(f"Last Down Event: {self.data['last_down_time']}")
            down_lbl.setObjectName("mutedLabel")
            layout.addWidget(down_lbl)

        timeline_title = QLabel("Recent Probing Timeline (Last 50):")
        timeline_title.setObjectName("strongLabel")
        layout.addWidget(timeline_title)

        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setMaximumHeight(200)
        logs = self.data.get("timeline_stats", [])
        self.log_text.setText("\n".join(logs[-50:]) if logs else "No timeline logs available.")
        layout.addWidget(self.log_text)

        btn_layout = QHBoxLayout()
        btn_layout.addStretch()
        close_btn = QPushButton("Close")
        close_btn.setObjectName("secondaryBtn")
        close_btn.setFixedHeight(34)
        close_btn.setMinimumWidth(100)
        close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(close_btn)
        layout.addLayout(btn_layout)

    def _apply_style(self):
        dark = self._dark
        bg = "#0a0d14" if dark else "#f8fafc"
        text = "#f8fafc" if dark else "#0f172a"
        muted = "#94a3b8" if dark else "#64748b"
        input_bg = "#0d121e" if dark else "#ffffff"
        border = "#1e293b" if dark else "#e2e8f0"
        self.setStyleSheet(f"""
            QDialog {{
                background-color: {bg};
                color: {text};
                font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            }}
            QLabel#detailTitle {{
                font-size: 17px;
                font-weight: 700;
                letter-spacing: -0.3px;
                color: {text};
                background: transparent;
            }}
            QLabel#mutedLabel {{
                color: {muted};
                font-size: 12px;
                font-weight: 500;
                background: transparent;
            }}
            QLabel#strongLabel {{
                color: {text};
                font-weight: 600;
                font-size: 13px;
                font-family: "JetBrains Mono", "Cascadia Code", "SF Mono", Consolas, monospace;
                background: transparent;
            }}
            QTextEdit {{
                background-color: {input_bg};
                color: {text};
                border: 1px solid {border};
                border-radius: 8px;
                padding: 12px;
                font-family: "JetBrains Mono", "Cascadia Code", "SF Mono", Consolas, monospace;
                font-size: 11px;
                line-height: 1.4;
            }}
            QPushButton#secondaryBtn {{
                background-color: {"#131a29" if dark else "#ffffff"};
                color: {text};
                border: 1px solid {border};
                border-radius: 7px;
                font-weight: 600;
                font-size: 12px;
                padding: 7px 18px;
            }}
            QPushButton#secondaryBtn:hover {{
                background-color: {"#1a2438" if dark else "#f1f5f9"};
                border-color: {"#2b3a56" if dark else "#cbd5e1"};
            }}
        """)


# ---------------------------------------------------------------------------
# DNS Issues Dialog
# ---------------------------------------------------------------------------
class DNSIssuesDialog(QDialog):
    """Inspector modal for hostnames that failed DNS resolution and were retired from probing."""

    def __init__(
        self,
        dns_issues: list[tuple[str, dict]],
        dark: bool = True,
        parent: QWidget | None = None,
    ):
        super().__init__(parent)
        self.dns_issues = list(dns_issues)
        self._dark = dark
        self.setWindowTitle("DNS Resolution Issues")
        self.resize(760, 480)
        self.setMinimumSize(600, 360)
        self._build_ui()
        self._apply_style()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(12)

        # Header section
        header_row = QHBoxLayout()
        header_row.setSpacing(10)

        badge = QLabel("DNS FAILURES")
        badge.setObjectName("badgeLabel")
        header_row.addWidget(badge)

        title = QLabel("Unresolved Hostnames")
        title.setObjectName("dnsDialogTitle")
        header_row.addWidget(title)
        header_row.addStretch()

        self.counter_label = QLabel(
            f"Total: {len(self.dns_issues)} host{'s' if len(self.dns_issues) != 1 else ''}"
        )
        self.counter_label.setObjectName("counterLabel")
        header_row.addWidget(self.counter_label)
        layout.addLayout(header_row)

        desc = QLabel(
            "The following target hostnames could not be resolved to an IP address. "
            "To conserve resources and avoid high-frequency lookup thrashing, they have been "
            "retired from the active probing queue and excluded from the Live Telemetry Chart."
        )
        desc.setWordWrap(True)
        desc.setObjectName("dnsDialogSubtitle")
        layout.addWidget(desc)

        # Search / filter toolbar
        search_row = QHBoxLayout()
        search_row.setSpacing(8)
        filter_lbl = QLabel("Filter:")
        filter_lbl.setObjectName("mutedLabel")
        search_row.addWidget(filter_lbl)

        self.search_box = QLineEdit()
        self.search_box.setPlaceholderText("Filter hostnames or error messages...")
        self.search_box.setClearButtonEnabled(True)
        self.search_box.textChanged.connect(self._filter_table)
        search_row.addWidget(self.search_box)
        layout.addLayout(search_row)

        # Issues Table
        self.table = QTableWidget()
        self.table.setColumnCount(4)
        self.table.setHorizontalHeaderLabels([
            "Target Hostname",
            "Status",
            "Error Details",
            "Detected At",
        ])
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        self.table.setShowGrid(True)

        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.Interactive)
        header.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.Stretch)
        header.setSectionResizeMode(3, QHeaderView.Interactive)
        self.table.setColumnWidth(0, 200)
        self.table.setColumnWidth(3, 150)

        self._populate_table()
        layout.addWidget(self.table, 1)

        # Footer Actions
        footer = QHBoxLayout()
        footer.setSpacing(10)

        self.feedback_label = QLabel("")
        self.feedback_label.setObjectName("feedbackLabel")
        footer.addWidget(self.feedback_label)
        footer.addStretch()

        self.copy_btn = QPushButton("Copy Hostnames")
        self.copy_btn.setObjectName("secondaryBtn")
        self.copy_btn.setToolTip("Copy hostnames to clipboard (selected or all filtered)")
        self.copy_btn.clicked.connect(self.copy_hostnames)
        footer.addWidget(self.copy_btn)

        self.export_btn = QPushButton("Export List...")
        self.export_btn.setObjectName("secondaryBtn")
        self.export_btn.setToolTip("Save unresolved hostnames report to CSV or TXT")
        self.export_btn.clicked.connect(self.export_issues)
        footer.addWidget(self.export_btn)

        close_btn = QPushButton("Close")
        close_btn.setObjectName("primaryActionBtn")
        close_btn.clicked.connect(self.accept)
        close_btn.setDefault(True)
        footer.addWidget(close_btn)

        layout.addLayout(footer)

    def _populate_table(self):
        self.table.setSortingEnabled(False)
        self.table.setRowCount(0)
        dark = self._dark
        text_color = QColor("#f8fafc" if dark else "#0f172a")
        mono_font = QFont("Consolas", 10)
        mono_font.setStyleHint(QFont.Monospace)

        for host, data in self.dns_issues:
            row = self.table.rowCount()
            self.table.insertRow(row)

            err_text = data.get("last_error") or "Name lookup failed"
            timestamp = data.get("last_down_time")
            if not timestamp or timestamp == "-":
                log_rows = data.get("log_rows")
                if log_rows:
                    timestamp = log_rows[-1][1]
                else:
                    timestamp = time.strftime("%Y-%m-%d %H:%M:%S")

            col_host = SortableTableWidgetItem(host)
            col_host.setForeground(text_color)
            col_host.setFont(mono_font)
            col_host.setTextAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            self.table.setItem(row, 0, col_host)

            col_status = SortableTableWidgetItem("DNS Error")
            col_status.setForeground(QColor("#fbbf24" if dark else "#b45309"))
            col_status.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(row, 1, col_status)

            col_err = SortableTableWidgetItem(err_text)
            col_err.setForeground(QColor("#94a3b8" if dark else "#64748b"))
            col_err.setToolTip(err_text)
            col_err.setTextAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            self.table.setItem(row, 2, col_err)

            col_time = SortableTableWidgetItem(str(timestamp))
            col_time.setForeground(text_color)
            col_time.setFont(mono_font)
            col_time.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(row, 3, col_time)

        self.table.setSortingEnabled(True)

    def _filter_table(self, query: str):
        query = query.strip().lower()
        visible_count = 0
        for row in range(self.table.rowCount()):
            host_item = self.table.item(row, 0)
            err_item = self.table.item(row, 2)
            matches = True
            if query:
                host_str = host_item.text().lower() if host_item else ""
                err_str = err_item.text().lower() if err_item else ""
                matches = query in host_str or query in err_str
            self.table.setRowHidden(row, not matches)
            if matches:
                visible_count += 1
        total = len(self.dns_issues)
        if query:
            self.counter_label.setText(f"Showing {visible_count} of {total} failed hosts")
        else:
            self.counter_label.setText(f"Total: {total} host{'s' if total != 1 else ''}")

    def copy_hostnames(self):
        selected_indexes = self.table.selectionModel().selectedRows()
        if selected_indexes:
            hosts = [
                self.table.item(idx.row(), 0).text()
                for idx in selected_indexes
                if self.table.item(idx.row(), 0)
            ]
            desc = f"{len(hosts)} selected hostname{'s' if len(hosts) != 1 else ''}"
        else:
            hosts = []
            for row in range(self.table.rowCount()):
                if not self.table.isRowHidden(row):
                    item = self.table.item(row, 0)
                    if item:
                        hosts.append(item.text())
            desc = f"{len(hosts)} hostname{'s' if len(hosts) != 1 else ''}"

        if not hosts:
            self.feedback_label.setText("No hostnames to copy.")
            self.feedback_label.setStyleSheet("color: #94a3b8; font-size: 11px;")
            return

        text = "\n".join(hosts)
        clipboard = QApplication.clipboard()
        if clipboard is not None:
            clipboard.setText(text)
            self.feedback_label.setText(f"✓ Copied {desc} to clipboard")
            self.feedback_label.setStyleSheet("color: #10b981; font-weight: 600; font-size: 11px;")
            QTimer.singleShot(4000, lambda: self.feedback_label.setText(""))

    def export_issues(self):
        if not self.dns_issues:
            QMessageBox.information(self, "Export", "No DNS issues to export.")
            return
        default_name = f"dns_issues_{time.strftime('%Y%m%d_%H%M%S')}.csv"
        path, sel_filter = QFileDialog.getSaveFileName(
            self,
            "Export DNS Issues",
            default_name,
            "CSV Files (*.csv);;Text Files (*.txt);;All Files (*)",
        )
        if not path:
            return
        try:
            is_csv = path.lower().endswith(".csv") or "CSV" in sel_filter
            with open(path, "w", newline="", encoding="utf-8-sig" if is_csv else "utf-8") as f:
                if is_csv:
                    writer = csv.writer(f)
                    writer.writerow([
                        "Target Hostname",
                        "Status",
                        "Error Details",
                        "Detected Timestamp",
                    ])
                    for row in range(self.table.rowCount()):
                        if not self.table.isRowHidden(row):
                            writer.writerow([
                                _sanitize_csv_cell(self.table.item(row, 0).text() if self.table.item(row, 0) else ""),
                                _sanitize_csv_cell(self.table.item(row, 1).text() if self.table.item(row, 1) else ""),
                                _sanitize_csv_cell(self.table.item(row, 2).text() if self.table.item(row, 2) else ""),
                                _sanitize_csv_cell(self.table.item(row, 3).text() if self.table.item(row, 3) else ""),
                            ])
                else:
                    f.write(
                        f"# DNS Resolution Failures Report - {time.strftime('%Y-%m-%d %H:%M:%S')}\n"
                    )
                    f.write(
                        f"# Total targets retired from probing: {len(self.dns_issues)}\n\n"
                    )
                    for row in range(self.table.rowCount()):
                        if not self.table.isRowHidden(row):
                            host = self.table.item(row, 0).text() if self.table.item(row, 0) else ""
                            err = self.table.item(row, 2).text() if self.table.item(row, 2) else ""
                            ts = self.table.item(row, 3).text() if self.table.item(row, 3) else ""
                            f.write(f"{host:<30} | {err:<40} | {ts}\n")
            self.feedback_label.setText(f"✓ Exported to {Path(path).name}")
            self.feedback_label.setStyleSheet("color: #10b981; font-weight: 600; font-size: 11px;")
            QTimer.singleShot(4000, lambda: self.feedback_label.setText(""))
        except Exception as exc:
            QMessageBox.critical(self, "Export Error", f"Failed to export DNS issues:\n{exc}")

    def _apply_style(self):
        dark = self._dark
        bg = "#0a0d14" if dark else "#f8fafc"
        text = "#f8fafc" if dark else "#0f172a"
        muted = "#94a3b8" if dark else "#64748b"
        border = "#1e293b" if dark else "#e2e8f0"
        header_bg = "#141b2c" if dark else "#f1f5f9"
        row_alt = "#0d121e" if dark else "#f8fafc"
        row_sel = "#1e293f" if dark else "#e2e8f0"
        input_bg = "#0d121e" if dark else "#ffffff"
        amber = "#fbbf24" if dark else "#b45309"
        amber_bg = "rgba(245, 158, 11, 0.12)" if dark else "#fef3c7"
        amber_border = "rgba(245, 158, 11, 0.4)" if dark else "#f59e0b"

        self.setStyleSheet(f"""
            QDialog {{
                background-color: {bg};
                color: {text};
                font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            }}
            QLabel#dnsDialogTitle {{
                font-size: 17px;
                font-weight: 700;
                letter-spacing: -0.3px;
                color: {text};
                background: transparent;
            }}
            QLabel#dnsDialogSubtitle {{
                color: {muted};
                font-size: 12px;
                line-height: 1.4;
                background: transparent;
            }}
            QLabel#counterLabel {{
                color: {muted};
                font-size: 12px;
                font-weight: 600;
                background: transparent;
            }}
            QLabel#mutedLabel {{
                color: {muted};
                font-size: 12px;
                font-weight: 500;
                background: transparent;
            }}
            QLabel#badgeLabel {{
                background-color: {amber_bg};
                color: {amber};
                border: 1px solid {amber_border};
                border-radius: 12px;
                padding: 3px 10px;
                font-size: 10px;
                font-weight: 700;
                letter-spacing: 0.6px;
            }}
            QLineEdit {{
                background-color: {input_bg};
                color: {text};
                border: 1px solid {border};
                border-radius: 7px;
                padding: 6px 11px;
                font-size: 12px;
            }}
            QLineEdit:focus {{
                border: 1px solid #38bdf8;
            }}
            QTableWidget {{
                background-color: {input_bg};
                alternate-background-color: {row_alt};
                color: {text};
                gridline-color: {border};
                border: 1px solid {border};
                border-radius: 8px;
                selection-background-color: {row_sel};
                selection-color: {text};
                font-size: 12px;
            }}
            QHeaderView::section {{
                background-color: {header_bg};
                color: {muted};
                font-weight: 700;
                font-size: 10.5px;
                border: none;
                border-bottom: 1px solid {border};
                padding: 8px 10px;
                text-transform: uppercase;
                letter-spacing: 0.6px;
            }}
            QPushButton#secondaryBtn {{
                background-color: {"#131a29" if dark else "#ffffff"};
                color: {text};
                border: 1px solid {border};
                border-radius: 7px;
                font-weight: 600;
                font-size: 12px;
                padding: 6px 14px;
            }}
            QPushButton#secondaryBtn:hover {{
                background-color: {"#1a2438" if dark else "#f1f5f9"};
                border-color: {"#2b3a56" if dark else "#cbd5e1"};
            }}
            QPushButton#primaryActionBtn {{
                background-color: #2563eb;
                color: #ffffff;
                border: 1px solid {"rgba(255,255,255,0.15)" if dark else "transparent"};
                border-radius: 7px;
                font-weight: 600;
                font-size: 12px;
                padding: 6px 18px;
            }}
            QPushButton#primaryActionBtn:hover {{
                background-color: #1d4ed8;
            }}
        """)


# ---------------------------------------------------------------------------
# Latency Chart
# ---------------------------------------------------------------------------
class LatencyChart(QWidget):
    COLORS_DARK = [
        "#38bdf8",  # Electric Cyan
        "#34d399",  # Emerald Mint
        "#fbbf24",  # Warm Amber
        "#f43f5e",  # Precision Crimson
        "#a78bfa",  # Soft Violet
        "#2dd4bf",  # Teal
        "#fb923c",  # Vivid Tangerine
        "#f472b6",  # Pink
        "#60a5fa",  # Royal Blue
        "#4ade80",  # Green
        "#e879f9",  # Fuchsia
        "#94a3b8",  # Steel Slate
    ]
    COLORS_LIGHT = [
        "#0284c7",  # Deep Sky
        "#059669",  # Emerald
        "#d97706",  # Warm Amber
        "#e11d48",  # Precision Crimson
        "#7c3aed",  # Violet
        "#0d9488",  # Teal
        "#ea580c",  # Vivid Tangerine
        "#db2777",  # Pink
        "#2563eb",  # Royal Blue
        "#16a34a",  # Forest
        "#c026d3",  # Fuchsia
        "#475569",  # Steel Slate
    ]

    def __init__(self, parent=None):
        super().__init__(parent)
        self.history = {}
        self.max_points = 60
        self._dark = True
        self.setMinimumHeight(240)

    def set_dark(self, dark: bool):
        self._dark = dark
        self.update()

    def set_data(self, host_data, max_points):
        self.max_points = max(1, int(max_points))
        self.history = {
            host: list(data["latencies"])[-self.max_points:]
            for host, data in host_data.items()
            if data.get("latencies")
        }
        self.update()

    @staticmethod
    def _elide(text, max_chars=16):
        text = str(text)
        return text if len(text) <= max_chars else text[: max_chars - 1] + "…"

    def paintEvent(self, _event):
        _event.accept()
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        dark = self._dark
        bg = QColor("#111624" if dark else "#ffffff")
        title_color = QColor("#f8fafc" if dark else "#0f172a")
        muted = QColor("#64748b" if dark else "#94a3b8")
        grid_color = QColor("#1c273c" if dark else "#f1f5f9")
        legend_color = QColor("#f8fafc" if dark else "#0f172a")
        legend_bg = QColor("#151d2f" if dark else "#f8fafc")
        colors = self.COLORS_DARK if dark else self.COLORS_LIGHT

        painter.fillRect(self.rect(), bg)
        painter.setPen(title_color)
        painter.setFont(QFont("Segoe UI", 9, QFont.Bold))
        painter.drawText(16, 22, "LIVE LATENCY TELEMETRY")

        n_hosts = len(self.history)
        legend_w = 144 if n_hosts else 0
        left_axis, top_m, bottom_m, right_gap = 52, 38, 30, 8
        plot = self.rect().adjusted(left_axis, top_m, -(legend_w + right_gap + 4), -bottom_m)
        legend_rect = QRectF(
            self.width() - legend_w - right_gap,
            top_m,
            legend_w,
            max(0, self.height() - top_m - bottom_m),
        )

        if not self.history:
            painter.setPen(muted)
            painter.setFont(QFont("Segoe UI", 10))
            painter.drawText(plot, Qt.AlignCenter, "Waiting for ICMP telemetry samples...")
            return

        valid = [v for vals in self.history.values() for v in vals if v > 0]
        max_latency = max(10.0, (max(valid) if valid else 10.0) * 1.15)

        grid = QPen(grid_color)
        painter.setPen(grid)
        for i in range(5):
            y = plot.top() + plot.height() * i / 4
            painter.drawLine(QPointF(plot.left(), y), QPointF(plot.right(), y))
            val = max_latency * (4 - i) / 4
            painter.setPen(muted)
            painter.setFont(QFont("Consolas", 8))
            painter.drawText(
                QRectF(2, y - 8, left_axis - 6, 16),
                Qt.AlignRight | Qt.AlignVCenter,
                f"{val:.0f}ms",
            )
            painter.setPen(grid)

        painter.setPen(QPen(grid_color))
        painter.drawRect(plot)
        painter.save()
        painter.setClipRect(plot)

        for idx, (host, values) in enumerate(self.history.items()):
            if not values:
                continue
            pen = QPen(QColor(colors[idx % len(colors)]))
            pen.setWidth(2)
            pen.setCapStyle(Qt.RoundCap)
            pen.setJoinStyle(Qt.RoundJoin)
            painter.setPen(pen)
            points = []
            n = len(values)
            for i, value in enumerate(values):
                x = plot.left() if n == 1 else plot.left() + plot.width() * i / (n - 1)
                y = plot.bottom() - plot.height() * min(value, max_latency) / max_latency
                points.append(QPointF(x, y))
            for i in range(1, len(points)):
                painter.drawLine(points[i - 1], points[i])
        painter.restore()

        painter.fillRect(legend_rect, legend_bg)
        painter.setPen(QPen(grid_color))
        painter.drawRoundedRect(legend_rect, 6, 6)

        row_h = 18
        max_rows = max(1, int(legend_rect.height() // row_h) - 1)
        items = list(self.history.items())
        shown = items[:max_rows]
        hidden = len(items) - len(shown)

        painter.setFont(QFont("Segoe UI", 8))
        for idx, item in enumerate(shown):
            host = item[0]
            ly = legend_rect.top() + 8 + idx * row_h
            color = QColor(colors[idx % len(colors)])
            lx = legend_rect.left() + 8
            painter.setBrush(QBrush(color))
            painter.setPen(Qt.NoPen)
            painter.drawRoundedRect(QRectF(lx, ly - 3, 10, 6), 3, 3)
            painter.setPen(legend_color)
            painter.drawText(
                QRectF(lx + 16, ly - 7, legend_w - 24, 14),
                Qt.AlignLeft | Qt.AlignVCenter,
                self._elide(host, 13),
            )

        if hidden > 0:
            ly = legend_rect.top() + 8 + len(shown) * row_h
            painter.setPen(muted)
            painter.drawText(
                QRectF(legend_rect.left() + 8, ly - 7, legend_w - 16, 14),
                Qt.AlignLeft | Qt.AlignVCenter,
                f"+{hidden} more...",
            )

        painter.setPen(muted)
        painter.setFont(QFont("Consolas", 8))
        painter.drawText(
            QRectF(plot.left(), plot.bottom() + 6, plot.width(), 18),
            Qt.AlignCenter,
            f"Buffer window: last {self.max_points} probes",
        )


# ---------------------------------------------------------------------------
# Card-like frame helper
# ---------------------------------------------------------------------------
class CardFrame(QFrame):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("cardFrame")
        self.setFrameShape(QFrame.NoFrame)


class MetricCard(QFrame):
    """Sleek telemetry KPI stat card for top ribbon."""

    def __init__(self, title: str, value: str = "—", subtitle: str = "", parent=None):
        super().__init__(parent)
        self.setObjectName("metricCard")
        self.setFrameShape(QFrame.NoFrame)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(3)

        self.lbl_title = QLabel(title.upper())
        self.lbl_title.setObjectName("metricTitle")
        self.lbl_title.setTextFormat(Qt.PlainText)

        self.lbl_val = QLabel(value)
        self.lbl_val.setObjectName("metricValue")
        self.lbl_val.setTextFormat(Qt.PlainText)

        self.lbl_sub = QLabel(subtitle)
        self.lbl_sub.setObjectName("metricSub")
        self.lbl_sub.setTextFormat(Qt.PlainText)

        layout.addWidget(self.lbl_title)
        layout.addWidget(self.lbl_val)
        layout.addWidget(self.lbl_sub)

    def set_data(self, value: str, subtitle: str | None = None):
        self.lbl_val.setText(str(value))
        if subtitle is not None:
            self.lbl_sub.setText(str(subtitle))


# ---------------------------------------------------------------------------
# Main Application
# ---------------------------------------------------------------------------
class PingMonitorApp(QMainWindow):
    PACKET_SIZE_BYTES = 56
    PING_INTERVAL_SECONDS = 1.0
    GRAPH_HISTORY_POINTS = 60
    TIMEOUT_SECONDS = 2.0
    MAX_LOG_ENTRIES = 1000
    CONFIG_FILE = Path(__file__).resolve().with_name("ping_monitor_config.json")

    def __init__(self):
        super().__init__()
        self.hosts = []
        self.config_warning = None
        self.is_monitoring = False
        self._run_id = 0
        self.probe_task = None
        self._stop_pending = False
        self.probe_pool = QThreadPool(self)
        self.probe_pool.setMaxThreadCount(1)
        self.host_data = {}
        self._host_items = {}
        self._last_down_signature = None
        self._dns_error_count = 0
        self._active_sort_col = None
        self._active_sort_order = Qt.AscendingOrder

        self.packet_size_bytes = self.PACKET_SIZE_BYTES
        self.ping_interval_seconds = self.PING_INTERVAL_SECONDS
        self.graph_history_points = self.GRAPH_HISTORY_POINTS
        self.timeout_seconds = self.TIMEOUT_SECONDS
        self.max_concurrent_probes = MAX_CONCURRENT_PROBES

        self.setWindowTitle("Ping Check — Professional Network Monitor")
        screen = QApplication.primaryScreen()
        if screen is not None:
            available = screen.availableGeometry()
            width = min(1260, int(available.width() * 0.92))
            height = min(880, int(available.height() * 0.9))
            self.resize(width, height)
            self.setMinimumSize(min(760, width), min(520, height))
        else:
            self.resize(1100, 740)
            self.setMinimumSize(760, 520)

        self._is_dark = True

        self.build_ui()
        self.apply_style()

        self.timer = QTimer(self)
        self.timer.timeout.connect(self.refresh_dashboard)
        self.timer.start(1000)

        self.load_config()

    # ------------------------------------------------------------------ UI
    def build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(16, 16, 16, 16)
        root.setSpacing(12)

        # Header bar
        header = QHBoxLayout()
        header.setContentsMargins(0, 0, 0, 0)
        
        title_col = QVBoxLayout()
        title_col.setSpacing(1)
        self.title_label = QLabel("Network Reachability Monitor")
        self.title_label.setObjectName("appTitle")
        self.subtitle_label = QLabel("Real-time ICMP telemetry & latency diagnostic tool")
        self.subtitle_label.setObjectName("appSubtitle")
        title_col.addWidget(self.title_label)
        title_col.addWidget(self.subtitle_label)
        header.addLayout(title_col, 1)

        right_box = QHBoxLayout()
        right_box.setSpacing(10)
        right_box.setAlignment(Qt.AlignRight | Qt.AlignVCenter)

        self.theme_btn = QPushButton()
        self.theme_btn.setFixedSize(36, 32)
        self.theme_btn.clicked.connect(self.toggle_theme)
        self.theme_btn.setObjectName("secondaryBtn")
        self.theme_btn.setCursor(Qt.PointingHandCursor)

        self.status_pill = QLabel("IDLE")
        self.status_pill.setObjectName("statusIdle")
        self.status_pill.setFixedHeight(32)
        self.status_pill.setAlignment(Qt.AlignCenter)
        self.status_pill.setMinimumWidth(116)

        right_box.addWidget(self.theme_btn)
        right_box.addWidget(self.status_pill)
        header.addLayout(right_box)
        root.addLayout(header)

        # Telemetry KPI Ribbon
        kpi_row = QHBoxLayout()
        kpi_row.setSpacing(12)
        self.card_fleet = MetricCard("Monitored Targets", "0 Targets", "Awaiting start")
        self.card_reachability = MetricCard("Fleet Reachability", "100%", "0 unreachable")
        self.card_latency = MetricCard("Average RTT", "—", "Min / Max —")
        self.card_loss = MetricCard("Global ICMP Loss", "0.0%", "0 dropped pkts")
        kpi_row.addWidget(self.card_fleet)
        kpi_row.addWidget(self.card_reachability)
        kpi_row.addWidget(self.card_latency)
        kpi_row.addWidget(self.card_loss)
        root.addLayout(kpi_row)

        # Target Configuration & Parameter Deck
        config_card = CardFrame()
        config_layout = QVBoxLayout(config_card)
        config_layout.setContentsMargins(16, 14, 16, 14)
        config_layout.setSpacing(10)

        section = QLabel("TARGET FLEET CONFIGURATION")
        section.setObjectName("targetConfigurationCaption")
        section.setTextFormat(Qt.PlainText)
        config_layout.addWidget(section)

        host_row = QHBoxLayout()
        host_row.setSpacing(8)
        self.host_entry = QLineEdit()
        self.host_entry.setText("8.8.8.8, 1.1.1.1, google.com")
        self.host_entry.setPlaceholderText("Enter target IPs or domains separated by commas (e.g. 8.8.8.8, 1.1.1.1, google.com)")
        self.host_entry.setClearButtonEnabled(True)
        self.host_entry.setFixedHeight(34)
        host_row.addWidget(self.host_entry, 1)

        self.import_btn = QPushButton("Import List")
        self.import_btn.clicked.connect(self.import_hosts_from_file)
        self.import_btn.setObjectName("secondaryBtn")
        self.import_btn.setFixedHeight(34)
        self.import_btn.setCursor(Qt.PointingHandCursor)
        host_row.addWidget(self.import_btn)

        self.save_config_btn = QPushButton("Save Config")
        self.save_config_btn.clicked.connect(self.save_config)
        self.save_config_btn.setObjectName("secondaryBtn")
        self.save_config_btn.setFixedHeight(34)
        self.save_config_btn.setCursor(Qt.PointingHandCursor)
        host_row.addWidget(self.save_config_btn)

        config_layout.addLayout(host_row)

        settings_row = QHBoxLayout()
        settings_row.setSpacing(12)

        pkt_col = QVBoxLayout()
        pkt_col.setSpacing(3)
        pkt_lbl = QLabel("PACKET SIZE")
        pkt_lbl.setObjectName("sectionCaption")
        pkt_lbl.setTextFormat(Qt.PlainText)
        pkt_col.addWidget(pkt_lbl)
        self.packet_spin = QSpinBox()
        self.packet_spin.setRange(0, 65500)
        self.packet_spin.setValue(56)
        self.packet_spin.setSuffix(" B")
        self.packet_spin.setMinimumWidth(100)
        self.packet_spin.setFixedHeight(34)
        self.packet_spin.setToolTip("ICMP payload size in bytes (0 to 65500).")
        pkt_col.addWidget(self.packet_spin)
        settings_row.addLayout(pkt_col)

        int_col = QVBoxLayout()
        int_col.setSpacing(3)
        int_lbl = QLabel("INTERVAL")
        int_lbl.setObjectName("sectionCaption")
        int_lbl.setTextFormat(Qt.PlainText)
        int_col.addWidget(int_lbl)
        self.interval_spin = QDoubleSpinBox()
        self.interval_spin.setRange(0.1, 3600.0)
        self.interval_spin.setDecimals(1)
        self.interval_spin.setSingleStep(0.1)
        self.interval_spin.setValue(1.0)
        self.interval_spin.setSuffix(" s")
        self.interval_spin.setMinimumWidth(100)
        self.interval_spin.setFixedHeight(34)
        self.interval_spin.setToolTip("Repeat cycle interval per host.")
        int_col.addWidget(self.interval_spin)
        settings_row.addLayout(int_col)

        timeout_col = QVBoxLayout()
        timeout_col.setSpacing(3)
        timeout_lbl = QLabel("TIMEOUT")
        timeout_lbl.setObjectName("sectionCaption")
        timeout_lbl.setTextFormat(Qt.PlainText)
        timeout_col.addWidget(timeout_lbl)
        self.timeout_spin = QDoubleSpinBox()
        self.timeout_spin.setRange(0.1, 10.0)
        self.timeout_spin.setDecimals(1)
        self.timeout_spin.setSingleStep(0.1)
        self.timeout_spin.setValue(2.0)
        self.timeout_spin.setSuffix(" s")
        self.timeout_spin.setMinimumWidth(100)
        self.timeout_spin.setFixedHeight(34)
        self.timeout_spin.setToolTip("Maximum wait time for each ICMP echo reply.")
        timeout_col.addWidget(self.timeout_spin)
        settings_row.addLayout(timeout_col)

        hist_col = QVBoxLayout()
        hist_col.setSpacing(3)
        hist_lbl = QLabel("WINDOW")
        hist_lbl.setObjectName("sectionCaption")
        hist_lbl.setTextFormat(Qt.PlainText)
        hist_col.addWidget(hist_lbl)
        self.history_spin = QSpinBox()
        self.history_spin.setRange(1, 10000)
        self.history_spin.setValue(60)
        self.history_spin.setSuffix(" pts")
        self.history_spin.setMinimumWidth(100)
        self.history_spin.setFixedHeight(34)
        self.history_spin.setToolTip("Graph telemetry sample window points.")
        hist_col.addWidget(self.history_spin)
        settings_row.addLayout(hist_col)

        concurrency_col = QVBoxLayout()
        concurrency_col.setSpacing(3)
        concurrency_lbl = QLabel("MAX PARALLEL")
        concurrency_lbl.setObjectName("sectionCaption")
        concurrency_lbl.setTextFormat(Qt.PlainText)
        concurrency_col.addWidget(concurrency_lbl)
        self.concurrency_spin = QSpinBox()
        self.concurrency_spin.setRange(1, MAX_CONCURRENT_PROBES)
        self.concurrency_spin.setValue(MAX_CONCURRENT_PROBES)
        self.concurrency_spin.setSuffix(" hosts")
        self.concurrency_spin.setMinimumWidth(110)
        self.concurrency_spin.setFixedHeight(34)
        self.concurrency_spin.setToolTip(
            f"Maximum number of simultaneous ICMP/DNS probes (application limit: {MAX_CONCURRENT_PROBES})."
        )
        concurrency_col.addWidget(self.concurrency_spin)
        settings_row.addLayout(concurrency_col)

        settings_row.addStretch()

        # Action Buttons aligned compactly
        btn_box = QHBoxLayout()
        btn_box.setSpacing(8)
        btn_box.setAlignment(Qt.AlignBottom)

        self.clear_btn = QPushButton("Clear")
        self.clear_btn.clicked.connect(self.clear_history)
        self.clear_btn.setEnabled(False)
        self.clear_btn.setObjectName("secondaryBtn")
        self.clear_btn.setFixedHeight(34)
        self.clear_btn.setCursor(Qt.PointingHandCursor)
        btn_box.addWidget(self.clear_btn)

        self.export_csv_btn = QPushButton("Export CSV")
        self.export_csv_btn.setEnabled(False)
        self.export_csv_btn.clicked.connect(self.export_report_csv)
        self.export_csv_btn.setObjectName("secondaryBtn")
        self.export_csv_btn.setFixedHeight(34)
        self.export_csv_btn.setCursor(Qt.PointingHandCursor)
        btn_box.addWidget(self.export_csv_btn)

        self.export_xlsx_btn = QPushButton("Export XLSX")
        self.export_xlsx_btn.setEnabled(False)
        self.export_xlsx_btn.clicked.connect(self.export_report_xlsx)
        self.export_xlsx_btn.setObjectName("secondaryBtn")
        self.export_xlsx_btn.setFixedHeight(34)
        self.export_xlsx_btn.setCursor(Qt.PointingHandCursor)
        btn_box.addWidget(self.export_xlsx_btn)

        self.export_txt_btn = QPushButton("Export TXT")
        self.export_txt_btn.setEnabled(False)
        self.export_txt_btn.clicked.connect(self.export_report_summary)
        self.export_txt_btn.setObjectName("secondaryBtn")
        self.export_txt_btn.setFixedHeight(34)
        self.export_txt_btn.setCursor(Qt.PointingHandCursor)
        btn_box.addWidget(self.export_txt_btn)

        self.start_btn = QPushButton("▶ Start Monitoring")
        self.start_btn.setObjectName("primaryBtn")
        self.start_btn.setFixedHeight(34)
        self.start_btn.setCursor(Qt.PointingHandCursor)
        self.start_btn.clicked.connect(self.toggle_monitoring)
        btn_box.addWidget(self.start_btn)

        settings_row.addLayout(btn_box)
        config_layout.addLayout(settings_row)
        root.addWidget(config_card)

        # Main splitter
        vertical = QSplitter(Qt.Vertical)
        root.addWidget(vertical, 1)

        # Upper Deck: Live Telemetry Table
        dash_card = CardFrame()
        dash_layout = QVBoxLayout(dash_card)
        dash_layout.setContentsMargins(14, 12, 14, 12)
        dash_layout.setSpacing(8)

        top = QHBoxLayout()
        top.setContentsMargins(0, 0, 0, 0)
        top.setSpacing(10)
        dash_title = QLabel("LIVE TELEMETRY TABLE")
        dash_title.setObjectName("strongLabel")
        dash_title.setTextFormat(Qt.PlainText)
        top.addWidget(dash_title)

        # Integrated Filter Input
        self.search_box = QLineEdit()
        self.search_box.setPlaceholderText("Filter targets...")
        self.search_box.textChanged.connect(self.filter_table)
        self.search_box.setMinimumWidth(180)
        self.search_box.setMaximumWidth(240)
        self.search_box.setFixedHeight(28)
        self.search_box.setObjectName("searchBox")
        top.addWidget(self.search_box)

        # Isolated DNS Failures Pill Button
        self.dns_issues_btn = QPushButton("0 DNS Issues")
        self.dns_issues_btn.setObjectName("dnsIssuesBtn")
        self.dns_issues_btn.setCursor(Qt.PointingHandCursor)
        self.dns_issues_btn.setFixedHeight(28)
        self.dns_issues_btn.setToolTip("View hostnames with DNS resolution failures (isolated from telemetry)")
        self.dns_issues_btn.clicked.connect(self.open_dns_issues_dialog)
        top.addWidget(self.dns_issues_btn)

        top.addStretch()

        self.host_count_label = QLabel("0 up / 0 total")
        self.host_count_label.setObjectName("captionLabel")
        self.host_count_label.setTextFormat(Qt.PlainText)
        top.addWidget(self.host_count_label)
        dash_layout.addLayout(top)

        self.table = QTableWidget()
        self.table.setObjectName("liveTelemetryTable")
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels([
            "Host / IP", "Status", "Seq", "Latency", "Packet Loss", "Failures",
        ])
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setAlternatingRowColors(False)
        self.table.setSortingEnabled(False)
        self.table.setShowGrid(True)
        self.table.doubleClicked.connect(self.on_table_double_click)
        h = self.table.horizontalHeader()
        h.setFixedHeight(32)
        h.setSectionsClickable(True)
        h.setSortIndicatorShown(False)
        h.sectionClicked.connect(self.on_table_header_clicked)
        h.setSectionResizeMode(QHeaderView.Stretch)
        h.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        dash_layout.addWidget(self.table)
        vertical.addWidget(dash_card)

        lower = QSplitter(Qt.Horizontal)

        chart_card = CardFrame()
        chart_layout = QVBoxLayout(chart_card)
        chart_layout.setContentsMargins(10, 10, 10, 10)
        self.chart = LatencyChart()
        chart_layout.addWidget(self.chart)
        lower.addWidget(chart_card)

        down_card = CardFrame()
        down_layout = QVBoxLayout(down_card)
        down_layout.setContentsMargins(14, 12, 14, 12)
        down_layout.setSpacing(6)
        down_title = QLabel("No ICMP Reply (Host May Still Be Up)")
        down_title.setObjectName("strongLabel")
        self.down_summary_label = QLabel("No targets monitored")
        self.down_summary_label.setObjectName("healthyLabel")
        down_layout.addWidget(down_title)
        down_layout.addWidget(self.down_summary_label)

        self.down_table = QTableWidget()
        self.down_table.setObjectName("noReplyTable")
        self.down_table.setColumnCount(5)
        self.down_table.setHorizontalHeaderLabels(
            ["Host", "Status", "Fails", "Loss", "Last Failed"]
        )
        self.down_table.verticalHeader().setVisible(False)
        self.down_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.down_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.down_table.setSortingEnabled(True)
        dh = self.down_table.horizontalHeader()
        dh.setFixedHeight(32)
        dh.setSectionResizeMode(QHeaderView.Interactive)
        dh.setStretchLastSection(True)
        self.down_table.setShowGrid(True)
        down_layout.addWidget(self.down_table)

        lower.addWidget(down_card)
        lower.setStretchFactor(0, 7)
        lower.setStretchFactor(1, 3)
        lower.setSizes([700, 300])
        lower.setChildrenCollapsible(False)
        vertical.addWidget(lower)
        vertical.setStretchFactor(0, 2)
        vertical.setStretchFactor(1, 3)

        self.footer_label = QLabel("ICMP Echo Prober • Ready")
        self.footer_label.setAlignment(Qt.AlignCenter)
        self.footer_label.setObjectName("appFooter")
        root.addWidget(self.footer_label)

    # ------------------------------------------------------------------ Theme
    def toggle_theme(self):
        self._is_dark = not self._is_dark
        self.apply_style()
        self.chart.set_dark(self._is_dark)
        self.chart.update()
        self.refresh_dashboard()

    def apply_style(self):
        dark = self._is_dark
        bg_main = "#0a0d14" if dark else "#f8fafc"
        bg_card = "#111624" if dark else "#ffffff"
        bg_input = "#0d121e" if dark else "#ffffff"
        border_input = "#1e293b" if dark else "#e2e8f0"
        border_focus = "#38bdf8" if dark else "#2563eb"
        text_main = "#f8fafc" if dark else "#0f172a"
        text_muted = "#94a3b8" if dark else "#64748b"
        text_caption = "#64748b" if dark else "#94a3b8"
        bg_table = "#0d121e" if dark else "#ffffff"
        grid_color = "#171f30" if dark else "#f1f5f9"
        header_bg = "#141b2c" if dark else "#f8fafc"
        header_text = "#94a3b8" if dark else "#475569"
        scroll_bg = "#0a0d14" if dark else "#f8fafc"
        scroll_handle = "#26334d" if dark else "#cbd5e1"
        scroll_hover = "#3b4e72" if dark else "#94a3b8"

        # Action-aware primary button
        if self.is_monitoring:
            primary = "#e11d48"
            primary_hover = "#be123c"
        else:
            primary = "#2563eb"
            primary_hover = "#1d4ed8"

        qss = f"""
        QMainWindow, QWidget {{
            background-color: {bg_main};
            color: {text_main};
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            font-size: 12px;
        }}
        QFrame#cardFrame {{
            background-color: {bg_card};
            color: {text_main};
            border: 1px solid {border_input};
            border-radius: 9px;
        }}
        QFrame#metricCard {{
            background-color: {bg_card};
            border: 1px solid {border_input};
            border-radius: 9px;
        }}
        QLabel#metricTitle {{
            color: {text_caption};
            font-size: 10px;
            font-weight: 700;
            letter-spacing: 0.8px;
            text-transform: uppercase;
        }}
        QLabel#metricValue {{
            color: {text_main};
            font-size: 21px;
            font-weight: 700;
            font-family: "JetBrains Mono", "Cascadia Code", "SF Mono", Consolas, monospace;
            letter-spacing: -0.5px;
        }}
        QLabel#metricSub {{
            color: {text_muted};
            font-size: 11px;
            font-weight: 500;
        }}
        QSplitter::handle {{
            background: {border_input};
            height: 2px;
        }}
        QLabel {{
            color: {text_main};
            background: transparent;
        }}
        QLabel#appTitle {{
            font-size: 20px;
            font-weight: 700;
            letter-spacing: -0.4px;
            color: {text_main};
        }}
        QLabel#appSubtitle {{
            color: {text_muted};
            font-size: 12px;
            font-weight: 400;
        }}
        QLabel#appFooter {{
            color: {text_caption};
            font-size: 11px;
            font-family: "JetBrains Mono", "Cascadia Code", "SF Mono", Consolas, monospace;
        }}
        QLabel#sectionCaption {{
            color: {text_muted};
            font-size: 10px;
            font-weight: 700;
            letter-spacing: 0.7px;
            text-transform: uppercase;
        }}
        QLabel#targetConfigurationCaption {{
            color: {text_muted};
            font-size: 10.5px;
            font-weight: 700;
            letter-spacing: 0.8px;
            text-transform: uppercase;
        }}
        QLabel#strongLabel {{
            color: {text_main};
            font-weight: 600;
            font-size: 13px;
        }}
        QLabel#bodyLabel {{
            color: {text_main};
            font-size: 12px;
        }}
        QLabel#captionLabel {{
            color: {text_caption};
            font-size: 11px;
        }}
        QLabel#healthyLabel {{
            color: {"#10b981" if dark else "#059669"};
            font-weight: 600;
        }}
        QLineEdit, QSpinBox, QDoubleSpinBox {{
            background-color: {bg_input};
            color: {text_main};
            border: 1px solid {border_input};
            border-radius: 7px;
            padding: 6px 11px;
            selection-background-color: {primary};
            selection-color: #ffffff;
            font-size: 12px;
        }}
        QLineEdit:focus, QSpinBox:focus, QDoubleSpinBox:focus {{
            border: 1px solid {border_focus};
        }}
        QPushButton#primaryBtn {{
            background-color: {primary};
            color: #ffffff;
            border: 1px solid {"rgba(255,255,255,0.15)" if dark else "transparent"};
            border-radius: 7px;
            font-weight: 600;
            font-size: 12px;
            letter-spacing: 0.2px;
            padding: 6px 18px;
        }}
        QPushButton#primaryBtn:hover {{
            background-color: {primary_hover};
        }}
        QPushButton#primaryBtn:disabled {{
            background-color: {"#182133" if dark else "#e2e8f0"};
            color: {"#475569" if dark else "#94a3b8"};
            border: 1px solid {"#1e293b" if dark else "#cbd5e1"};
        }}
        QPushButton#secondaryBtn {{
            background-color: {"#131a29" if dark else "#ffffff"};
            color: {text_main};
            border: 1px solid {border_input};
            border-radius: 7px;
            font-weight: 500;
            font-size: 12px;
            padding: 6px 14px;
        }}
        QPushButton#secondaryBtn:hover {{
            background-color: {"#1a2438" if dark else "#f1f5f9"};
            border-color: {"#2b3a56" if dark else "#cbd5e1"};
        }}
        QPushButton#secondaryBtn:disabled {{
            color: {text_caption};
            border-color: {border_input};
            background-color: {"#0d131f" if dark else "#f8fafc"};
        }}
        QTableWidget {{
            background-color: {bg_table};
            color: {text_main};
            gridline-color: {grid_color};
            border: 1px solid {border_input};
            border-radius: 8px;
            outline: none;
            font-size: 12px;
        }}
        QTableWidget::item {{
            color: {text_main};
            background-color: {bg_table};
            padding: 6px 9px;
            border: none;
        }}
        QTableWidget::item:selected {{
            background-color: {"#1a2538" if dark else "#e2e8f0"};
            color: {text_main};
        }}
        QHeaderView::section {{
            background-color: {header_bg};
            color: {header_text};
            border: none;
            border-bottom: 1px solid {border_input};
            padding: 8px 10px;
            font-size: 10px;
            font-weight: 700;
            letter-spacing: 0.6px;
            text-transform: uppercase;
        }}
        QHeaderView {{
            background-color: {header_bg};
        }}
        QScrollBar:vertical, QScrollBar:horizontal {{
            background: {scroll_bg};
            border: none;
            width: 7px;
            height: 7px;
        }}
        QScrollBar::handle:vertical, QScrollBar::handle:horizontal {{
            background: {scroll_handle};
            border-radius: 3px;
        }}
        QScrollBar::handle:vertical:hover, QScrollBar::handle:horizontal:hover {{
            background: {scroll_hover};
        }}
        QTableWidget#liveTelemetryTable QScrollBar:vertical, QTableWidget#noReplyTable QScrollBar:vertical {{
            width: 8px;
            background: transparent;
            margin: 32px 2px 2px 2px;
        }}
        QTableWidget#liveTelemetryTable QScrollBar::handle:vertical, QTableWidget#noReplyTable QScrollBar::handle:vertical {{
            min-height: 40px;
            background: {scroll_handle};
            border-radius: 4px;
        }}
        QTableWidget#liveTelemetryTable QScrollBar::handle:vertical:hover, QTableWidget#noReplyTable QScrollBar::handle:vertical:hover {{
            background: {scroll_hover};
        }}
        QScrollBar::add-line, QScrollBar::sub-line {{
            height: 0;
            width: 0;
        }}
        QToolTip {{
            background-color: {bg_card};
            color: {text_main};
            border: 1px solid {border_input};
            padding: 5px 9px;
            border-radius: 6px;
            font-size: 11px;
        }}
        """
        self.setStyleSheet(qss)
        theme_name = "Dark" if dark else "Light"
        self.theme_btn.setText("☾" if dark else "☀")
        self.theme_btn.setToolTip(f"{theme_name} theme selected. Click to switch.")
        self.theme_btn.setAccessibleName(f"{theme_name} theme selected")

        if not self.is_monitoring:
            if self._stop_pending:
                self._set_status_stopping()
            else:
                self._set_status_idle()
        else:
            self._set_status_running()

        if self.host_data:
            down_hosts = [
                (h, d) for h, d in self.host_data.items()
                if self._is_down_status(d.get("last_status"))
            ]
            if not down_hosts:
                self.down_summary_label.setStyleSheet(
                    f"color: {'#10b981' if dark else '#059669'}; font-weight: 600; background: transparent;"
                )
            else:
                self.down_summary_label.setStyleSheet(
                    f"color: {'#f43f5e' if dark else '#e11d48'}; font-weight: 600; background: transparent;"
                )

        self._refresh_table_row_themes()
        if hasattr(self, "dns_issues_btn"):
            self.update_dns_pill(getattr(self, "_dns_error_count", 0))

    def _refresh_table_row_themes(self):
        dark = self._is_dark
        text_color = QColor("#f8fafc" if dark else "#0f172a")
        bg_color = QColor("#0d121e" if dark else "#ffffff")
        for table in (self.table, self.down_table):
            for row in range(table.rowCount()):
                for col in range(table.columnCount()):
                    item = table.item(row, col)
                    if item is not None:
                        item.setBackground(bg_color)
                        if col == 1:
                            status = item.text()
                            if status == "Connected":
                                item.setForeground(QColor("#10b981" if dark else "#059669"))
                            elif status == "No ICMP Reply":
                                item.setForeground(QColor("#fbbf24" if dark else "#d97706"))
                            elif status in ("Queued", "Idle"):
                                item.setForeground(QColor("#94a3b8" if dark else "#64748b"))
                            else:
                                item.setForeground(QColor("#f43f5e" if dark else "#e11d48"))
                        else:
                            item.setForeground(text_color)

    def _set_status_idle(self):
        dark = self._is_dark
        bg = "#121826" if dark else "#f1f5f9"
        border = "#1f2a3f" if dark else "#e2e8f0"
        text = "#94a3b8" if dark else "#64748b"
        self.status_pill.setText("IDLE")
        self.status_pill.setStyleSheet(
            f"background: {bg}; color: {text}; border: 1px solid {border}; "
            "border-radius: 14px; font-weight: 700; font-size: 11px; letter-spacing: 0.8px; padding: 2px 10px;"
        )

    def _set_status_running(self):
        dark = self._is_dark
        bg = "rgba(16, 185, 129, 0.12)" if dark else "#ecfdf5"
        border = "rgba(16, 185, 129, 0.4)" if dark else "#a7f3d0"
        text = "#34d399" if dark else "#059669"
        self.status_pill.setText("MONITORING")
        self.status_pill.setStyleSheet(
            f"background: {bg}; color: {text}; border: 1px solid {border}; "
            "border-radius: 14px; font-weight: 700; font-size: 11px; letter-spacing: 0.8px; padding: 2px 10px;"
        )

    def _set_status_stopping(self):
        dark = self._is_dark
        bg = "rgba(245, 158, 11, 0.12)" if dark else "#fffbeb"
        border = "rgba(245, 158, 11, 0.4)" if dark else "#fde68a"
        text = "#fbbf24" if dark else "#d97706"
        self.status_pill.setText("STOPPING")
        self.status_pill.setStyleSheet(
            f"background: {bg}; color: {text}; border: 1px solid {border}; "
            "border-radius: 14px; font-weight: 700; font-size: 11px; letter-spacing: 0.8px; padding: 2px 10px;"
        )

    # ------------------------------------------------------------------ Helpers
    @staticmethod
    def _is_down_status(status):
        return status in ("No ICMP Reply", "Unreachable")

    @staticmethod
    def _compute_host_stats(
        data: dict,
    ) -> tuple[int, int, float | None, float | None, float | None, float | None]:
        sent = data["sent"]
        received = data["received"]
        loss_pct = round(((sent - received) / sent) * 100, 1) if sent > 0 else None
        latency_count = data.get("latency_count", 0)
        if latency_count:
            min_lat = data["min_latency"]
            max_lat = data["max_latency"]
            avg_lat = round(data["latency_sum"] / latency_count, 2)
        else:
            min_lat = max_lat = avg_lat = None
        return sent, received, loss_pct, min_lat, max_lat, avg_lat

    def _open_file_dialog(self, title, name_filter, save=False, default_name="", default_suffix=""):
        dialog = QFileDialog(self, title)
        dialog.setOption(QFileDialog.Option.DontUseNativeDialog, True)
        dialog.setViewMode(QFileDialog.ViewMode.Detail)
        dialog.setNameFilter(name_filter)
        dialog.setOption(QFileDialog.Option.ReadOnly, not save)
        start_dir = os.path.dirname(os.path.abspath(__file__))
        if os.path.isdir(start_dir):
            dialog.setDirectory(start_dir)
        if save:
            dialog.setAcceptMode(QFileDialog.AcceptMode.AcceptSave)
            dialog.setFileMode(QFileDialog.FileMode.AnyFile)
            if default_suffix:
                dialog.setDefaultSuffix(default_suffix)
            if default_name:
                dialog.selectFile(default_name)
        else:
            dialog.setAcceptMode(QFileDialog.AcceptMode.AcceptOpen)
            dialog.setFileMode(QFileDialog.FileMode.ExistingFile)
        if dialog.exec():
            selected = dialog.selectedFiles()
            if selected:
                path = selected[0]
                if default_suffix and not path.lower().endswith("." + default_suffix.lower()):
                    path = path + "." + default_suffix
                return path
        return ""

    def _msg(self, title, text, icon=QMessageBox.Information):
        box = QMessageBox(self)
        box.setWindowTitle(title)
        box.setText(text)
        box.setIcon(icon)
        box.setStandardButtons(QMessageBox.Ok)
        dark = self._is_dark
        bg = "#0d1117" if dark else "#ffffff"
        text_c = "#e6edf3" if dark else "#1f2328"
        box.setStyleSheet(f"""
            QMessageBox {{ background-color: {bg}; color: {text_c}; }}
            QLabel {{ color: {text_c}; }}
            QPushButton {{
                background-color: {"#238636" if dark else "#1f883d"}; color: white;
                border: none; border-radius: 6px; padding: 5px 14px; min-width: 60px;
            }}
        """)
        box.exec()

    def _confirm(self, title, text):
        box = QMessageBox(self)
        box.setWindowTitle(title)
        box.setText(text)
        box.setIcon(QMessageBox.Question)
        box.setStandardButtons(QMessageBox.Yes | QMessageBox.No)
        box.setDefaultButton(QMessageBox.No)
        dark = self._is_dark
        bg = "#0d1117" if dark else "#ffffff"
        text_c = "#e6edf3" if dark else "#1f2328"
        box.setStyleSheet(f"""
            QMessageBox {{ background-color: {bg}; color: {text_c}; }}
            QLabel {{ color: {text_c}; }}
            QPushButton {{
                background-color: {"#238636" if dark else "#1f883d"}; color: white;
                border: none; border-radius: 6px; padding: 5px 14px; min-width: 60px;
            }}
        """)
        return box.exec() == QMessageBox.Yes

    # ------------------------------------------------------------------ Import / Config
    def import_hosts_from_file(self):
        file_path = self._open_file_dialog(
            "Import Host List",
            "Host Lists (*.txt *.csv);;Text Files (*.txt);;CSV Files (*.csv);;All Files (*)",
            save=False,
        )
        if not file_path:
            return

        try:
            hosts = []
            seen = set()

            def add_host(value):
                value = str(value).strip().strip("\"'")
                if not value:
                    return
                if value.lower() in {"host", "hostname", "ip", "ip address", "address", "host/ip"}:
                    return

                for segment in re.split(r"[,;\t]+", value):
                    segment = segment.strip().strip("\"'")
                    paren = re.search(r"\(([^()]+)\)", segment)
                    if paren:
                        segment = paren.group(1).strip()
                    for candidate in re.split(r"\s+", segment):
                        candidate = candidate.strip().strip("\"'")
                        if not candidate or candidate.lower() in {"host", "hostname", "ip", "address"}:
                            continue
                        normalized = candidate.casefold()
                        if normalized not in seen:
                            seen.add(normalized)
                            hosts.append(candidate)

            if file_path.lower().endswith(".csv"):
                with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
                    rows = list(csv.reader(f))
                if rows:
                    headers = [c.strip().lower() for c in rows[0]]
                    ip_cols = [
                        i
                        for i, h in enumerate(headers)
                        if h in {"ip", "ip address", "ip_address", "address", "host", "hostname", "host/ip"}
                        or "ip" in h
                    ]
                    data_rows = rows[1:] if any(headers) else rows
                    if ip_cols:
                        preferred = next(
                            (i for i, h in enumerate(headers) if h in {"ip", "ip address", "ip_address"}),
                            ip_cols[0],
                        )
                        for row in data_rows:
                            if preferred < len(row):
                                add_host(row[preferred])
                            elif row:
                                add_host(row[0])
                    else:
                        for row in rows:
                            for cell in row:
                                add_host(cell)
            else:
                with open(file_path, "r", encoding="utf-8-sig") as f:
                    for line in f:
                        line = line.strip()
                        if line:
                            add_host(line)

            if not hosts:
                self._msg("No Hosts Found", "No usable hosts found in file.", QMessageBox.Warning)
                return

            current = [h.strip() for h in self.host_entry.text().split(",") if h.strip()]
            merged = []
            seen_merged = set()
            for host in current + hosts:
                if host not in seen_merged:
                    seen_merged.add(host)
                    merged.append(host)
            self.host_entry.setText(", ".join(merged))
            self._msg("Hosts Imported", f"Successfully imported {len(hosts)} host(s).")
        except Exception as e:
            self._msg("Import Error", f"Failed to import file:\n{e}", QMessageBox.Critical)

    def save_config(self):
        config = {
            "hosts": self.host_entry.text(),
            "packet_size": self.packet_spin.value(),
            "interval": self.interval_spin.value(),
            "timeout": self.timeout_spin.value(),
            "history_points": self.history_spin.value(),
            "max_concurrent": self.concurrency_spin.value(),
            "dark_theme": self._is_dark,
        }
        config_path = Path(self.CONFIG_FILE)
        temporary_path = None
        try:
            config_path.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(
                "w",
                encoding="utf-8",
                dir=config_path.parent,
                prefix=f".{config_path.name}.",
                suffix=".tmp",
                delete=False,
            ) as config_file:
                json.dump(config, config_file, indent=2)
                config_file.flush()
                os.fsync(config_file.fileno())
                temporary_path = Path(config_file.name)
            os.replace(temporary_path, config_path)
            temporary_path = None
            self._msg("Config Saved", f"Settings saved to {config_path}")
        except (OSError, TypeError, ValueError) as error:
            self._msg("Error", f"Failed to save config:\n{error}", QMessageBox.Critical)
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    def load_config(self):
        config_path = Path(self.CONFIG_FILE)
        if not config_path.exists():
            return
        try:
            config = json.loads(config_path.read_text(encoding="utf-8"))
            if not isinstance(config, dict):
                raise ValueError("The config root must be a JSON object.")

            hosts = config.get("hosts")
            if hosts is not None:
                if not isinstance(hosts, str):
                    raise ValueError("The 'hosts' setting must be text.")
            pending_values = {}
            for key, control, integer_only in (
                ("packet_size", self.packet_spin, True),
                ("interval", self.interval_spin, False),
                ("timeout", self.timeout_spin, False),
                ("history_points", self.history_spin, True),
                ("max_concurrent", self.concurrency_spin, True),
            ):
                if key not in config:
                    continue
                value = config[key]
                if isinstance(value, bool) or not isinstance(value, (int, float)):
                    raise ValueError(f"The '{key}' setting must be numeric.")
                if integer_only and not isinstance(value, int):
                    raise ValueError(f"The '{key}' setting must be an integer.")
                if not math.isfinite(value):
                    raise ValueError(f"The '{key}' setting must be finite.")
                minimum = control.minimum()
                maximum = control.maximum()
                if value < minimum or value > maximum:
                    raise ValueError(
                        f"The '{key}' setting must be between {minimum} and {maximum}."
                    )
                pending_values[key] = value

            dark_theme = self._is_dark
            if "dark_theme" in config:
                if not isinstance(config["dark_theme"], bool):
                    raise ValueError("The 'dark_theme' setting must be true or false.")
                dark_theme = config["dark_theme"]

            # Apply only after every field passes validation, so malformed files
            # leave the complete set of defaults intact.
            if hosts is not None:
                self.host_entry.setText(hosts)
            for key, value in pending_values.items():
                {
                    "packet_size": self.packet_spin,
                    "interval": self.interval_spin,
                    "timeout": self.timeout_spin,
                    "history_points": self.history_spin,
                    "max_concurrent": self.concurrency_spin,
                }[key].setValue(value)
            self._is_dark = dark_theme

            self.apply_style()
            self.chart.set_dark(self._is_dark)
            self.chart.update()
        except (OSError, json.JSONDecodeError, OverflowError, TypeError, ValueError) as error:
            self.config_warning = str(error)

    # ------------------------------------------------------------------ Monitoring
    def toggle_monitoring(self):
        if self.is_monitoring:
            self.stop_monitoring()
            return

        try:
            self.hosts, duplicate_count = normalize_targets(self.host_entry.text())
        except ValueError as error:
            self._msg("Invalid Targets", str(error), QMessageBox.Warning)
            return

        if not self.hosts:
            self._msg("No Hosts", "Please enter at least one target host.", QMessageBox.Warning)
            return

        self.host_entry.setText(", ".join(self.hosts))
        if duplicate_count:
            self._msg(
                "Duplicate Targets Removed",
                f"Monitoring each target once; removed {duplicate_count} duplicate entr{'y' if duplicate_count == 1 else 'ies'}.",
                QMessageBox.Information,
            )

        self.packet_size_bytes = self.packet_spin.value()
        self.ping_interval_seconds = self.interval_spin.value()
        self.graph_history_points = self.history_spin.value()
        self.timeout_seconds = self.timeout_spin.value()
        self.max_concurrent_probes = self.concurrency_spin.value()

        self.host_data.clear()
        self._host_items.clear()
        self._last_down_signature = None
        self.table.setRowCount(0)
        self.down_table.setRowCount(0)
        self.update_dns_pill(0)

        self._run_id += 1
        for host in self.hosts:
            self.host_data[host] = {
                "latencies": deque(maxlen=self.graph_history_points),
                "latency_count": 0,
                "latency_sum": 0.0,
                "min_latency": None,
                "max_latency": None,
                "sent": 0,
                "checks": 0,
                "received": 0,
                "timeline_stats": [],
                "log_rows": [],
                "tree_id": host,
                "last_status": None,
                "last_error": None,
                "consecutive_fails": 0,
                "last_down_time": "-",
            }
            row = self.table.rowCount()
            self.table.insertRow(row)
            values = [host, "Queued", "0", "-", "0%", "0"]
            sort_vals = {2: 0, 3: float("inf"), 4: 0.0, 5: 0}
            dark = self._is_dark
            text_color = QColor("#f8fafc" if dark else "#0f172a")
            bg_color = QColor("#0d121e" if dark else "#ffffff")
            mono_font = QFont("Consolas", 10)
            mono_font.setStyleHint(QFont.Monospace)
            for col, value in enumerate(values):
                item = SortableTableWidgetItem(str(value), numeric=col in (2, 3, 4, 5), sort_value=sort_vals.get(col))
                item.setTextAlignment(Qt.AlignCenter)
                item.setForeground(text_color)
                item.setBackground(bg_color)
                if col in (2, 3, 4, 5):
                    item.setFont(mono_font)
                self.table.setItem(row, col, item)
                if col == 0:
                    self._host_items[host] = item

        self.is_monitoring = True
        self._stop_pending = False
        self.host_count_label.setText(f"0 up / {len(self.host_data)} total")
        self.set_controls_enabled(False)
        self.start_btn.setText("⏹ Stop Monitoring")
        self.apply_style()
        self._set_status_running()
        self.export_txt_btn.setEnabled(False)
        self.export_csv_btn.setEnabled(False)
        self.export_xlsx_btn.setEnabled(False)
        self.clear_btn.setEnabled(False)
        task = ProbeTask(
            run_id=self._run_id,
            targets=self.hosts,
            payload_size=self.packet_size_bytes,
            timeout=self.timeout_seconds,
            interval=self.ping_interval_seconds,
            max_concurrent=self.max_concurrent_probes,
        )
        if hasattr(task.signals, "batch_completed"):
            task.signals.batch_completed.connect(self._handle_probe_batch)
        task.signals.completed.connect(self._handle_probe_result)
        task.signals.finished.connect(self._handle_probe_task_finished)
        self.probe_task = task
        self.probe_pool.start(task)

    def stop_monitoring(self):
        if not self.is_monitoring:
            return
        self.is_monitoring = False
        self._stop_pending = True
        if self.probe_task is not None:
            self.probe_task.request_stop()
            self.set_controls_enabled(False)
            self.start_btn.setText("Stopping…")
            self.start_btn.setEnabled(False)
            self._set_status_stopping()
            self.footer_label.setText("Finishing current probes…")
        else:
            self._finalize_stop()

    def _finalize_stop(self):
        self._stop_pending = False
        self.probe_task = None
        self.set_controls_enabled(True)
        self.start_btn.setEnabled(True)
        self.start_btn.setText("▶ Start Monitoring")
        self.apply_style()
        self._set_status_idle()
        self.export_txt_btn.setEnabled(bool(self.host_data))
        self.export_csv_btn.setEnabled(bool(self.host_data))
        self.export_xlsx_btn.setEnabled(bool(self.host_data))
        self.clear_btn.setEnabled(bool(self.host_data))
        self.update_footer_stats()
        if self.host_data and self._confirm(
            "Export Report", "Monitoring stopped. Export performance report summary?"
        ):
            self.export_report_summary()

    @Slot(int)
    def _handle_probe_task_finished(self, run_id):
        if run_id != self._run_id:
            return
        self.probe_task = None
        if self._stop_pending:
            self._finalize_stop()
        elif self.is_monitoring:
            self.is_monitoring = False
            self._finalize_stop()
            self._msg(
                "Probe Runner Stopped",
                "The background probe runner stopped unexpectedly. You can start monitoring again.",
                QMessageBox.Warning,
            )

    def _process_single_result(self, target: str, result: ProbeResult, update_table: bool = True, trigger_sort: bool = False):
        if target not in self.host_data:
            return
        if not isinstance(result, ProbeResult):
            result = ProbeResult("Probe Error", error="The probe returned an invalid result.")

        data = self.host_data[target]
        valid_statuses = {
            "Connected", "No ICMP Reply", "DNS Error", "Permission Denied", "Probe Error", "Unreachable"
        }
        if result.status not in valid_statuses:
            result = ProbeResult("Probe Error", error=f"Unknown probe status: {result.status!r}")
        elif result.status == "Connected" and (
            result.latency_ms is None
            or not isinstance(result.latency_ms, (int, float))
            or not math.isfinite(result.latency_ms)
            or result.latency_ms < 0
        ):
            result = ProbeResult("Probe Error", error="The probe returned an invalid latency value.")
        if not isinstance(result.packet_sent, bool):
            result = ProbeResult("Probe Error", error="The probe returned invalid send metadata.", packet_sent=False)
        elif result.status == "Connected" and not result.packet_sent:
            result = ProbeResult("Probe Error", error="A connected result must correspond to a sent packet.", packet_sent=False)
        elif result.status in {"DNS Error", "Permission Denied", "Probe Error"} and result.packet_sent:
            result = ProbeResult(
                result.status, error=result.error, packet_sent=False
            )
        data["checks"] = data.get("checks", data["sent"]) + 1
        if result.packet_sent:
            data["sent"] += 1
        sequence = data["checks"]
        timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
        status = result.status
        data["last_status"] = status
        data["last_error"] = result.error

        if status == "Connected" and result.latency_ms is not None and result.packet_sent:
            latency = result.latency_ms
            data["received"] += 1
            data["latencies"].append(latency)
            data["latency_count"] += 1
            data["latency_sum"] += latency
            data["min_latency"] = latency if data["min_latency"] is None else min(data["min_latency"], latency)
            data["max_latency"] = latency if data["max_latency"] is None else max(data["max_latency"], latency)
            data["consecutive_fails"] = 0
            latency_text = f"{latency:.2f} ms"
            timeline = f"[{timestamp}] Seq {sequence}: Success — Latency {latency:.2f} ms"
            log_result = (sequence, timestamp, "Success", latency)
        else:
            if result.packet_sent:
                data["latencies"].append(0)
            data["consecutive_fails"] += 1
            if self._is_down_status(status):
                data["last_down_time"] = timestamp
            latency_text = "—"
            error_text = f" — {result.error}" if result.error else ""
            timeline = f"[{timestamp}] Seq {sequence}: FAILED — {status}{error_text}"
            log_result = (sequence, timestamp, status, result.error or "")

        data["timeline_stats"].append(timeline)
        data["log_rows"].append(log_result)
        del data["timeline_stats"][:-self.MAX_LOG_ENTRIES]
        del data["log_rows"][:-self.MAX_LOG_ENTRIES]

        if update_table:
            if status == "DNS Error":
                self.remove_host_row(target)
            else:
                loss_pct = (
                    round(((data["sent"] - data["received"]) / data["sent"]) * 100, 1)
                    if data["sent"] else None
                )
                row_values = (
                    target,
                    status,
                    str(sequence),
                    latency_text,
                    f"{loss_pct}%" if loss_pct is not None else "n/a",
                    str(data["consecutive_fails"]),
                )
                sort_values = {
                    2: sequence,
                    3: latency if (status == "Connected" and latency is not None) else float("inf"),
                    4: loss_pct if loss_pct is not None else float("-1"),
                    5: data["consecutive_fails"],
                }
                self.update_row(target, row_values, trigger_sort=trigger_sort, sort_values=sort_values)

    @Slot(int, str, object)
    def _handle_probe_result(self, run_id, target, result):
        if run_id != self._run_id or not self.is_monitoring:
            return
        # If batching is active, ignore redundant individual completed signal
        if hasattr(self, "probe_task") and self.probe_task is not None and hasattr(self.probe_task.signals, "batch_completed"):
            return
        self._process_single_result(target, result, update_table=True, trigger_sort=False)
        self.update_footer_stats()

    @Slot(int, list)
    def _handle_probe_batch(self, run_id, batch):
        if run_id != self._run_id or not self.is_monitoring:
            return
        if not batch:
            return
        for item in batch:
            if isinstance(item, (tuple, list)) and len(item) == 2:
                self._process_single_result(item[0], item[1], update_table=True, trigger_sort=False)
        self.update_footer_stats()

    def set_controls_enabled(self, enabled):
        self.host_entry.setEnabled(enabled)
        self.import_btn.setEnabled(enabled)
        self.save_config_btn.setEnabled(enabled)
        self.packet_spin.setEnabled(enabled)
        self.interval_spin.setEnabled(enabled)
        self.timeout_spin.setEnabled(enabled)
        self.history_spin.setEnabled(enabled)
        self.concurrency_spin.setEnabled(enabled)

    def _is_host_in_table(self, host: str) -> bool:
        if host in self._host_items:
            return True
        for r in range(self.table.rowCount()):
            it = self.table.item(r, 0)
            if it and it.text() == host:
                return True
        return False

    @Slot(str)
    def remove_host_row(self, host: str):
        col0_item = self._host_items.pop(host, None)
        if col0_item is not None:
            row = col0_item.row()
            if row >= 0:
                self.table.removeRow(row)
                return
        for r in range(self.table.rowCount()):
            it = self.table.item(r, 0)
            if it and it.text() == host:
                self.table.removeRow(r)
                break

    def on_table_header_clicked(self, col: int):
        header = self.table.horizontalHeader()
        if self._active_sort_col == col:
            order = (
                Qt.DescendingOrder
                if self._active_sort_order == Qt.AscendingOrder
                else Qt.AscendingOrder
            )
        else:
            order = Qt.DescendingOrder if col in (3, 4, 5) else Qt.AscendingOrder
        self._active_sort_col = col
        self._active_sort_order = order
        header.setSortIndicatorShown(True)
        header.setSortIndicator(col, order)
        self.table.sortItems(col, order)

    @Slot(str, object)
    def update_row(self, host, values, trigger_sort: bool = False, sort_values: dict | None = None):
        col0_item = self._host_items.get(host)
        if col0_item is None:
            return
        row = col0_item.row()
        if row < 0:
            return

        dark = self._is_dark
        text_color = QColor("#f8fafc" if dark else "#0f172a")
        bg_color = QColor("#0d121e" if dark else "#ffffff")
        mono_font = QFont("Consolas", 10)
        mono_font.setStyleHint(QFont.Monospace)

        for col, value in enumerate(values):
            val_str = str(value)
            sv = sort_values.get(col) if sort_values else None
            item = self.table.item(row, col)
            if item is None:
                item = SortableTableWidgetItem(val_str, numeric=col in (2, 3, 4, 5), sort_value=sv)
                item.setTextAlignment(Qt.AlignCenter)
                item.setForeground(text_color)
                item.setBackground(bg_color)
                if col in (2, 3, 4, 5):
                    item.setFont(mono_font)
                self.table.setItem(row, col, item)
                if col == 0:
                    self._host_items[host] = item
            else:
                if hasattr(item, "set_value"):
                    item.set_value(val_str, sort_value=sv)
                else:
                    if item.text() != val_str:
                        item.setText(val_str)
                    if sv is not None:
                        item.setData(Qt.UserRole, sv)

        status = str(values[1])
        status_item = self.table.item(row, 1)
        if status_item is not None:
            if status == "Connected":
                status_item.setForeground(QColor("#10b981" if dark else "#059669"))
            elif status == "No ICMP Reply":
                status_item.setForeground(QColor("#fbbf24" if dark else "#d97706"))
            elif status in ("Queued", "Idle"):
                status_item.setForeground(QColor("#94a3b8" if dark else "#64748b"))
            else:
                status_item.setForeground(QColor("#f43f5e" if dark else "#e11d48"))
            status_item.setToolTip(self.host_data.get(host, {}).get("last_error") or "")

    def filter_table(self, text):
        text = text.strip().lower()
        for row in range(self.table.rowCount()):
            item = self.table.item(row, 0)
            if item:
                self.table.setRowHidden(row, text not in item.text().lower())

    def on_table_double_click(self, index):
        row = index.row()
        host_item = self.table.item(row, 0)
        if host_item:
            host = host_item.text()
            data = self.host_data.get(host)
            if data:
                dlg = HostDetailDialog(host, data, dark=self._is_dark, parent=self)
                dlg.exec()

    def update_dns_pill(self, count: int):
        self._dns_error_count = count
        if not hasattr(self, "dns_issues_btn"):
            return
        dark = self._is_dark
        if count == 0:
            self.dns_issues_btn.setText("0 DNS Issues")
            if dark:
                self.dns_issues_btn.setStyleSheet("""
                    QPushButton#dnsIssuesBtn {
                        background-color: #111624;
                        color: #64748b;
                        border: 1px solid #1e293b;
                        border-radius: 14px;
                        padding: 3px 12px;
                        font-size: 11px;
                        font-weight: 500;
                    }
                    QPushButton#dnsIssuesBtn:hover {
                        background-color: #151d2f;
                        color: #94a3b8;
                        border-color: #2b3a56;
                    }
                """)
            else:
                self.dns_issues_btn.setStyleSheet("""
                    QPushButton#dnsIssuesBtn {
                        background-color: #f1f5f9;
                        color: #64748b;
                        border: 1px solid #e2e8f0;
                        border-radius: 14px;
                        padding: 3px 12px;
                        font-size: 11px;
                        font-weight: 500;
                    }
                    QPushButton#dnsIssuesBtn:hover {
                        background-color: #e2e8f0;
                        color: #334155;
                        border-color: #cbd5e1;
                    }
                """)
        else:
            self.dns_issues_btn.setText(f"⚠ {count} DNS Issue{'s' if count != 1 else ''}")
            if dark:
                self.dns_issues_btn.setStyleSheet("""
                    QPushButton#dnsIssuesBtn {
                        background-color: rgba(245, 158, 11, 0.12);
                        color: #fbbf24;
                        border: 1px solid rgba(245, 158, 11, 0.4);
                        border-radius: 14px;
                        padding: 3px 12px;
                        font-size: 11px;
                        font-weight: 700;
                    }
                    QPushButton#dnsIssuesBtn:hover {
                        background-color: rgba(245, 158, 11, 0.22);
                        border: 1px solid #fbbf24;
                    }
                """)
            else:
                self.dns_issues_btn.setStyleSheet("""
                    QPushButton#dnsIssuesBtn {
                        background-color: #fef3c7;
                        color: #b45309;
                        border: 1px solid #f59e0b;
                        border-radius: 14px;
                        padding: 3px 12px;
                        font-size: 11px;
                        font-weight: 700;
                    }
                    QPushButton#dnsIssuesBtn:hover {
                        background-color: #fde68a;
                        border: 1px solid #d97706;
                    }
                """)

    def open_dns_issues_dialog(self):
        dns_issues = [
            (host, data)
            for host, data in self.host_data.items()
            if data.get("last_status") == "DNS Error"
        ]
        dlg = DNSIssuesDialog(dns_issues, dark=self._is_dark, parent=self)
        dlg.exec()

    def clear_history(self):
        if self.is_monitoring or not self.host_data:
            return
        if not self._confirm("Clear History", "Remove all monitoring data and history?"):
            return
        self.host_data.clear()
        self._host_items.clear()
        self._last_down_signature = None
        self._active_sort_col = None
        self._active_sort_order = Qt.AscendingOrder
        self.table.horizontalHeader().setSortIndicatorShown(False)
        self.table.setRowCount(0)
        self.down_table.setRowCount(0)
        self.host_count_label.setText("0 up / 0 total")
        self.down_summary_label.setText("No targets monitored")
        dark = self._is_dark
        self.down_summary_label.setStyleSheet(
            f"color: {'#3fb950' if dark else '#1a7f37'}; font-weight: 600; background: transparent;"
        )
        self.chart.set_data({}, self.graph_history_points)
        self.update_dns_pill(0)
        self.export_txt_btn.setEnabled(False)
        self.export_csv_btn.setEnabled(False)
        self.export_xlsx_btn.setEnabled(False)
        self.clear_btn.setEnabled(False)
        self.update_footer_stats()

    def refresh_dashboard(self):
        if not self.host_data:
            self.update_footer_stats()
            self.update_dns_pill(0)
            return

        dns_issue_count = sum(
            1 for data in self.host_data.values()
            if data.get("last_status") == "DNS Error"
        )
        if dns_issue_count > 0:
            for host, data in list(self.host_data.items()):
                if data.get("last_status") == "DNS Error" and (host in self._host_items or self._is_host_in_table(host)):
                    self.remove_host_row(host)

        # Periodic controlled table re-sort (at most once per second, never on raw packet updates)
        if self._active_sort_col is not None:
            self.table.sortItems(self._active_sort_col, self._active_sort_order)

        up_count = sum(data.get("last_status") == "Connected" for data in self.host_data.values())
        monitored_count = len(self.host_data) - dns_issue_count
        if dns_issue_count > 0:
            self.host_count_label.setText(f"{up_count} up / {monitored_count} monitored")
        else:
            self.host_count_label.setText(f"{up_count} up / {len(self.host_data)} total")

        # Live telemetry chart: strictly only include pingable ("Connected") and timeouts ("No ICMP Reply")
        # Excludes DNS resolution failures and unhandled probe exceptions
        chart_targets = {
            host: data
            for host, data in self.host_data.items()
            if data.get("last_status") in ("Connected", "No ICMP Reply")
        }
        self.chart.set_data(chart_targets, self.graph_history_points)

        self.update_dns_pill(dns_issue_count)
        pending_count = sum(
            1 for data in self.host_data.values()
            if data.get("checks", data.get("sent", 0)) == 0
        )
        down_hosts = [
            (host, data)
            for host, data in self.host_data.items()
            if self._is_down_status(data.get("last_status"))
        ]
        probe_issue_signature = tuple(
            (host, data.get("last_status"))
            for host, data in self.host_data.items()
            if data.get("last_status") in {"DNS Error", "Permission Denied", "Probe Error"}
        )
        signature = (
            tuple(
                (
                    host,
                    data.get("last_status"),
                    data.get("consecutive_fails", 0),
                    data.get("last_down_time", "-"),
                )
                for host, data in down_hosts
            ),
            pending_count,
            probe_issue_signature,
        )
        if signature != self._last_down_signature:
            self._last_down_signature = signature
            self.down_table.setSortingEnabled(False)
            self.down_table.setRowCount(0)
            dark = self._is_dark
            probe_issues = len(probe_issue_signature)
            summary_parts = []
            if down_hosts:
                summary_parts.append(f"{len(down_hosts)} of {len(self.host_data)} down")
            if probe_issues:
                summary_parts.append(f"{probe_issues} DNS/probe issue(s)")
            if pending_count:
                summary_parts.append(f"{pending_count} awaiting first result")

            if summary_parts:
                self.down_summary_label.setText(" · ".join(summary_parts))
                if down_hosts:
                    color = "#f43f5e" if dark else "#e11d48"
                elif probe_issues:
                    color = "#f59e0b" if dark else "#d97706"
                else:
                    color = "#848d97" if dark else "#656d76"
                self.down_summary_label.setStyleSheet(
                    f"color: {color}; font-weight: 600; background: transparent;"
                )
            else:
                self.down_summary_label.setText("No current ICMP failures")
                self.down_summary_label.setStyleSheet(
                    f"color: {'#10b981' if dark else '#059669'}; font-weight: 600; background: transparent;"
                )
            if down_hosts:
                text_color = QColor("#f8fafc" if dark else "#0f172a")
                bg_color = QColor("#0d121e" if dark else "#ffffff")
                mono_font = QFont("Consolas", 10)
                mono_font.setStyleHint(QFont.Monospace)
                for host, data in down_hosts:
                    row = self.down_table.rowCount()
                    self.down_table.insertRow(row)
                    sent = data["sent"]
                    received = data["received"]
                    loss = round(((sent - received) / sent) * 100, 1) if sent else None
                    values = [
                        host,
                        data.get("last_status", "Down"),
                        data.get("consecutive_fails", 0),
                        f"{loss}%" if loss is not None else "n/a",
                        data.get("last_down_time", "-"),
                    ]
                    for col, value in enumerate(values):
                        item = SortableTableWidgetItem(str(value), numeric=col in (2, 3))
                        item.setTextAlignment(Qt.AlignCenter)
                        item.setForeground(text_color)
                        item.setBackground(bg_color)
                        if col in (2, 3):
                            item.setFont(mono_font)
                        self.down_table.setItem(row, col, item)
                    status_col_item = self.down_table.item(row, 1)
                    if status_col_item is not None:
                        status_val = str(values[1])
                        if status_val == "No ICMP Reply":
                            status_col_item.setForeground(QColor("#fbbf24" if dark else "#d97706"))
                        else:
                            status_col_item.setForeground(QColor("#f43f5e" if dark else "#e11d48"))
            self.down_table.setSortingEnabled(True)
        self.update_footer_stats()

    def update_footer_stats(self):
        if not self.host_data:
            self.footer_label.setText("No targets yet • Ready")
            self.card_fleet.set_data("0 Targets", "Awaiting start")
            self.card_reachability.set_data("100%", "0 unreachable")
            self.card_latency.set_data("—", "Min / Max —")
            self.card_loss.set_data("0.0%", "0 dropped pkts")
            return

        total_hosts = len(self.host_data)
        dns_issue_count = sum(d.get("last_status") == "DNS Error" for d in self.host_data.values())
        monitored_hosts = total_hosts - dns_issue_count
        up_count = sum(data.get("last_status") == "Connected" for data in self.host_data.values())
        down_count = sum(self._is_down_status(d.get("last_status")) for d in self.host_data.values())
        reachability_pct = round((up_count / monitored_hosts * 100), 1) if monitored_hosts else 0.0

        total_sent = sum(d["sent"] for d in self.host_data.values())
        total_recv = sum(d["received"] for d in self.host_data.values())
        overall_loss = ((total_sent - total_recv) / total_sent * 100) if total_sent else None

        latency_sum = sum(d.get("latency_sum", 0.0) for d in self.host_data.values())
        latency_cnt = sum(d.get("latency_count", 0) for d in self.host_data.values())
        if latency_cnt:
            avg_lat_str = f"{latency_sum / latency_cnt:.2f} ms"
            minima = [d["min_latency"] for d in self.host_data.values() if d.get("min_latency") is not None]
            maxima = [d["max_latency"] for d in self.host_data.values() if d.get("max_latency") is not None]
            lat_sub = f"Min {min(minima):.1f} / Max {max(maxima):.1f} ms" if minima and maxima else f"{latency_cnt} samples"
        else:
            avg_lat_str = "—"
            lat_sub = "Min / Max —"

        if dns_issue_count > 0:
            self.card_fleet.set_data(f"{monitored_hosts} Monitored", f"{up_count} online / {dns_issue_count} DNS failures")
        else:
            self.card_fleet.set_data(f"{total_hosts} Targets", f"{up_count} online / {down_count} down")
        self.card_reachability.set_data(f"{reachability_pct}%", f"{up_count} reachable")
        self.card_latency.set_data(avg_lat_str, lat_sub)
        loss_str = f"{overall_loss:.1f}%" if overall_loss is not None else "0.0%"
        self.card_loss.set_data(loss_str, f"{total_sent - total_recv} dropped packets")

        no_reply = sum(
            self._is_down_status(d.get("last_status")) for d in self.host_data.values()
        )
        probe_issues = sum(
            d.get("last_status") in {"DNS Error", "Permission Denied", "Probe Error"}
            for d in self.host_data.values()
        )
        state = "Waiting for first results" if self.is_monitoring and not total_sent else (
            "Monitoring" if self.is_monitoring else "Stopped"
        )
        self.footer_label.setText(
            f"{len(self.host_data)} targets  |  ICMP loss: {overall_loss:.1f}%" if overall_loss is not None else
            f"{len(self.host_data)} targets  |  ICMP loss: n/a"
        )
        self.footer_label.setText(
            self.footer_label.text()
            + f"  |  No reply: {no_reply}  |  DNS/probe issues: {probe_issues}  |  {state}"
        )

    # ------------------------------------------------------------------ Export
    def _summary_info(self, data_snapshot):
        statuses = [data.get("last_status") for data in data_snapshot.values()]
        total_sent = sum(data.get("sent", 0) for data in data_snapshot.values())
        total_received = sum(data.get("received", 0) for data in data_snapshot.values())
        latency_samples = sum(data.get("latency_count", 0) for data in data_snapshot.values())
        latency_sum = sum(data.get("latency_sum", 0.0) for data in data_snapshot.values())
        minima = [data.get("min_latency") for data in data_snapshot.values()
                  if data.get("min_latency") is not None]
        maxima = [data.get("max_latency") for data in data_snapshot.values()
                  if data.get("max_latency") is not None]
        loss_pct = (
            round((total_sent - total_received) / total_sent * 100, 1)
            if total_sent else None
        )
        return {
            "Report generated": time.strftime("%Y-%m-%d %H:%M:%S"),
            "Monitoring state": "Monitoring" if self.is_monitoring else "Stopped",
            "Target status counts based on": "Each target's latest probe result",
            "Total targets": len(data_snapshot),
            "Targets responding to ICMP": sum(status == "Connected" for status in statuses),
            "Targets with no ICMP reply": sum(self._is_down_status(status) for status in statuses),
            "DNS errors": sum(status == "DNS Error" for status in statuses),
            "Permission errors": sum(status == "Permission Denied" for status in statuses),
            "Probe errors": sum(status == "Probe Error" for status in statuses),
            "Awaiting first result": sum(status is None for status in statuses),
            "Completed checks": sum(
                data.get("checks", data.get("sent", 0)) for data in data_snapshot.values()
            ),
            "ICMP requests sent": total_sent,
            "ICMP replies received": total_received,
            "ICMP packet loss (%)": loss_pct if loss_pct is not None else "N/A",
            "Latency samples": latency_samples,
            "Minimum latency (ms)": min(minima) if minima else "N/A",
            "Average latency (ms)": round(latency_sum / latency_samples, 2) if latency_samples else "N/A",
            "Maximum latency (ms)": max(maxima) if maxima else "N/A",
        }

    @staticmethod
    def _style_workbook_table(sheet, header_row=1):
        from openpyxl.styles import Alignment, Font, PatternFill

        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.sheet_view.showGridLines = False
        if sheet.max_row >= header_row and sheet.max_column:
            sheet.auto_filter.ref = (
                f"A{header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
            )
        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")

    def export_report_summary(self):
        if not self.host_data:
            self._msg("No Data", "No monitoring data available for export.", QMessageBox.Warning)
            return

        file_path = self._open_file_dialog(
            "Export Summary Report",
            "Text Files (*.txt);;All Files (*)",
            save=True,
            default_name=f"network_report_{time.strftime('%Y%m%d_%H%M%S')}.txt",
            default_suffix="txt",
        )
        if not file_path:
            return

        data_snapshot = {host: dict(data) for host, data in self.host_data.items()}
        summary_info = self._summary_info(data_snapshot)

        try:
            with open(file_path, "w", encoding="utf-8") as f:
                f.write("=" * 70 + "\n")
                f.write("             NETWORK PERFORMANCE SUMMARY REPORT\n")
                f.write(f"             Generated: {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write("=" * 70 + "\n\n")

                f.write("[1. SUMMARY INFO]\n")
                f.write("-" * 50 + "\n")
                for metric, value in summary_info.items():
                    f.write(f"{metric:<34}: {value}\n")
                f.write("\n")

                f.write("[2. HOST AGGREGATE METRICS]\n")
                f.write("-" * 80 + "\n")
                f.write(
                    f"{'Target Host':<20} | {'Sent':<6} | {'Recv':<6} | {'Loss %':<8} | "
                    f"{'Min':<9} | {'Max':<9} | {'Avg':<9}\n"
                )
                f.write("-" * 80 + "\n")

                for host, data in data_snapshot.items():
                    sent, received, loss_pct, min_lat, max_lat, avg_lat = self._compute_host_stats(data)
                    if min_lat is None:
                        min_lat, max_lat, avg_lat = "-", "-", "-"
                    else:
                        min_lat, max_lat, avg_lat = f"{min_lat}ms", f"{max_lat}ms", f"{avg_lat}ms"
                    loss_text = f"{loss_pct}%" if loss_pct is not None else "n/a"
                    f.write(
                        f"{host:<20} | {sent:<6} | {received:<6} | {loss_text:<8} | "
                        f"{min_lat:<9} | {max_lat:<9} | {avg_lat:<9}\n"
                    )
                f.write("-" * 80 + "\n\n")

                f.write("[3. TARGETS WITHOUT AN ICMP REPLY]\n")
                f.write("-" * 50 + "\n")
                down_now = [
                    h for h, d in data_snapshot.items() if self._is_down_status(d.get("last_status"))
                ]
                if down_now:
                    for h in down_now:
                        f.write(
                            f"• {h} (last failed: {data_snapshot[h].get('last_down_time', '-')})\n"
                        )
                else:
                    f.write("No target recorded an ICMP timeout. Web-service health is not tested.\n")
                f.write("\n")

                f.write("[4. DETAILED SEQUENCE TIMELINES]\n")
                for host, data in data_snapshot.items():
                    f.write(f"\n# Stream: {host}\n")
                    f.write("-" * 50 + "\n")
                    for log_entry in data["timeline_stats"]:
                        f.write(log_entry + "\n")

            self._msg("Success", f"Report successfully exported to:\n{file_path}")
        except Exception as e:
            self._msg("Error", f"Failed to write report:\n{str(e)}", QMessageBox.Critical)

    def export_report_xlsx(self):
        if not self.host_data:
            self._msg("No Data", "No monitoring data available for export.", QMessageBox.Warning)
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            self._msg(
                "XLSX Export Unavailable",
                "Install the project dependencies to enable Excel export:\n"
                "python -m pip install -r requirements.txt",
                QMessageBox.Warning,
            )
            return

        file_path = self._open_file_dialog(
            "Export Excel Workbook",
            "Excel Workbooks (*.xlsx);;All Files (*)",
            save=True,
            default_name=f"network_report_{time.strftime('%Y%m%d_%H%M%S')}.xlsx",
            default_suffix="xlsx",
        )
        if not file_path:
            return

        data_snapshot = {host: dict(data) for host, data in self.host_data.items()}
        summary_info = self._summary_info(data_snapshot)
        try:
            workbook = Workbook()
            summary_sheet = workbook.active
            summary_sheet.title = "Summary Info"
            summary_sheet.merge_cells("A1:B1")
            summary_sheet["A1"] = "Ping Monitor Summary Info"
            summary_sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
            summary_sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
            summary_sheet["A1"].alignment = Alignment(vertical="center")
            summary_sheet.row_dimensions[1].height = 28
            summary_sheet.append(["Generated", summary_info["Report generated"]])
            summary_sheet.append(["Monitoring state", summary_info["Monitoring state"]])
            summary_sheet.append([])
            summary_sheet.append(["Metric", "Value"])
            for metric, value in summary_info.items():
                if metric not in {"Report generated", "Monitoring state"}:
                    summary_sheet.append([metric, value])
            self._style_workbook_table(summary_sheet, header_row=4)
            summary_sheet.column_dimensions["A"].width = 34
            summary_sheet.column_dimensions["B"].width = 24

            host_sheet = workbook.create_sheet("Host Details")
            host_sheet.append([
                "Host", "Last Status", "Checks", "ICMP Sent", "ICMP Received",
                "Loss %", "Min ms", "Max ms", "Avg ms", "Consecutive Failures",
                "Last No Reply",
            ])
            for host, data in data_snapshot.items():
                sent, received, loss, minimum, maximum, average = self._compute_host_stats(data)
                host_sheet.append([
                    host, data.get("last_status") or "Awaiting first result",
                    data.get("checks", sent), sent, received,
                    loss if loss is not None else "N/A",
                    minimum if minimum is not None else "N/A",
                    maximum if maximum is not None else "N/A",
                    average if average is not None else "N/A",
                    data.get("consecutive_fails", 0), data.get("last_down_time", "-"),
                ])
            self._style_workbook_table(host_sheet)
            host_sheet.column_dimensions["A"].width = 30
            host_sheet.column_dimensions["B"].width = 24
            for column in "CDEFGHIJ":
                host_sheet.column_dimensions[column].width = 18
            host_sheet.column_dimensions["K"].width = 22

            probe_sheet = workbook.create_sheet("Probe Details")
            probe_sheet.append(["Host", "Sequence", "Timestamp", "Status", "Latency ms", "Details"])
            for host, data in data_snapshot.items():
                for sequence, timestamp, status, details in data.get("log_rows", []):
                    probe_sheet.append([
                        host, sequence, timestamp,
                        "Connected" if status == "Success" else status,
                        details if status == "Success" else None,
                        "" if status == "Success" else details,
                    ])
            self._style_workbook_table(probe_sheet)
            for column, width in {"A": 30, "B": 12, "C": 22, "D": 24, "E": 16, "F": 50}.items():
                probe_sheet.column_dimensions[column].width = width

            workbook.save(file_path)
            self._msg("Success", f"Excel workbook successfully exported to:\n{file_path}")
        except Exception as e:
            self._msg("Error", f"Failed to write Excel workbook:\n{e}", QMessageBox.Critical)

    def export_report_csv(self):
        if not self.host_data:
            self._msg("No Data", "No monitoring data available for export.", QMessageBox.Warning)
            return

        file_path = self._open_file_dialog(
            "Export CSV Report",
            "CSV Files (*.csv);;All Files (*)",
            save=True,
            default_name=f"network_report_{time.strftime('%Y%m%d_%H%M%S')}.csv",
            default_suffix="csv",
        )
        if not file_path:
            return

        data_snapshot = {host: dict(data) for host, data in self.host_data.items()}

        try:
            with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
                fieldnames = [
                    "Section", "Host", "Sequence", "Status", "Sent", "Received",
                    "Loss_Pct", "Latency_ms", "Min_ms", "Max_ms", "Avg_ms",
                    "Consecutive_Fails", "Last_Down_At", "Timestamp", "Details",
                ]
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                def write_sanitized_row(row_dict):
                    writer.writerow({k: _sanitize_csv_cell(v) for k, v in row_dict.items()})

                generated = time.strftime("%Y-%m-%d %H:%M:%S")
                for host, data in data_snapshot.items():
                    sent, received, loss_pct, min_lat, max_lat, avg_lat = self._compute_host_stats(data)
                    write_sanitized_row({
                        "Section": "Summary",
                        "Host": host,
                        "Status": data.get("last_status") or "-",
                        "Sent": sent,
                        "Received": received,
                        "Loss_Pct": loss_pct if loss_pct is not None else "",
                        "Min_ms": min_lat if min_lat is not None else "",
                        "Max_ms": max_lat if max_lat is not None else "",
                        "Avg_ms": avg_lat if avg_lat is not None else "",
                        "Consecutive_Fails": data.get("consecutive_fails", 0),
                        "Last_Down_At": data.get("last_down_time", "-"),
                        "Timestamp": generated,
                    })
                    if self._is_down_status(data.get("last_status")):
                        write_sanitized_row({
                            "Section": "No ICMP Reply",
                            "Host": host,
                            "Status": data.get("last_status"),
                            "Consecutive_Fails": data.get("consecutive_fails", 0),
                            "Last_Down_At": data.get("last_down_time", "-"),
                            "Details": data.get("last_error") or "",
                        })
                    if data.get("last_status") == "DNS Error":
                        write_sanitized_row({
                            "Section": "DNS Failures",
                            "Host": host,
                            "Status": "DNS Error",
                            "Consecutive_Fails": data.get("consecutive_fails", 0),
                            "Last_Down_At": data.get("last_down_time", "-"),
                            "Timestamp": generated,
                            "Details": data.get("last_error") or "Name resolution failed",
                        })
                    for sequence, timestamp, status, details in data.get("log_rows", []):
                        write_sanitized_row({
                            "Section": "Probe",
                            "Host": host,
                            "Sequence": sequence,
                            "Status": status,
                            "Latency_ms": details if status == "Success" else "",
                            "Timestamp": timestamp,
                            "Details": "" if status == "Success" else details,
                        })
            self._msg("Success", f"CSV successfully exported to:\n{file_path}")
        except Exception as e:
            self._msg("Error", f"Failed to write CSV:\n{str(e)}", QMessageBox.Critical)

    def closeEvent(self, event):
        self.is_monitoring = False
        self._run_id += 1
        self.timer.stop()
        if self.probe_task is not None:
            self.probe_task.request_stop()
            self.probe_task = None
        self.probe_pool.clear()
        self.probe_pool.waitForDone(int((self.timeout_seconds + 2) * 1000))
        event.accept()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    app = QApplication(sys.argv)
    app.setApplicationName("Ping Check Enhanced")
    app.setFont(QFont("Segoe UI", 9))

    window = PingMonitorApp()
    window.show()
    if window.config_warning:
        QTimer.singleShot(
            0,
            lambda warning=window.config_warning: window._msg(
                "Config Not Loaded", f"Using defaults because the config could not be loaded:\n{warning}",
                QMessageBox.Warning,
            ),
        )
    try:
        exit_code = app.exec()
    except KeyboardInterrupt:
        exit_code = 0
    sys.exit(exit_code)


if __name__ == "__main__":
    main()
