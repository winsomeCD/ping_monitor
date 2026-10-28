import csv
import json
import os
import re
import sys
import tempfile
import unittest
from collections import deque
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from openpyxl import load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR))

from PySide6.QtCore import QEventLoop, QTimer, Qt
from PySide6.QtWidgets import QApplication, QTableWidgetItem

import ping_monitor
from probes import ProbeResult


class PingMonitorAppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])
        ping_monitor._ensure_system_fonts()

    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.original_config_path = ping_monitor.PingMonitorApp.CONFIG_FILE
        ping_monitor.PingMonitorApp.CONFIG_FILE = Path(self.temporary_directory.name) / "config.json"
        self.window = ping_monitor.PingMonitorApp()
        self.window._msg = lambda *args: None
        self.window._confirm = lambda *args: False

    def tearDown(self):
        self.window.close()
        self.application.processEvents()
        ping_monitor.PingMonitorApp.CONFIG_FILE = self.original_config_path
        self.temporary_directory.cleanup()

    def test_theme_is_saved_and_restored(self):
        self.assertEqual(self.window.theme_btn.text(), "☾")
        self.window.toggle_theme()
        self.assertEqual(self.window.theme_btn.text(), "☀")
        self.window.save_config()

        saved_config = json.loads(Path(self.window.CONFIG_FILE).read_text(encoding="utf-8"))
        self.assertIs(saved_config["dark_theme"], False)

        restored_window = ping_monitor.PingMonitorApp()
        self.assertFalse(restored_window._is_dark)
        self.assertEqual(restored_window.theme_btn.text(), "☀")
        restored_window.close()

    def test_invalid_config_is_reported(self):
        Path(self.window.CONFIG_FILE).write_text('{"interval": "fast"}', encoding="utf-8")

        self.window.load_config()

        self.assertIn("interval", self.window.config_warning)

    def test_invalid_config_is_rejected_without_partially_applying_values(self):
        self.window.host_entry.setText("keep.example")
        Path(self.window.CONFIG_FILE).write_text(
            '{"hosts": "replace.example", "packet_size": 1200, "interval": 99999}',
            encoding="utf-8",
        )

        self.window.load_config()

        self.assertIn("interval", self.window.config_warning)
        self.assertEqual(self.window.host_entry.text(), "keep.example")
        self.assertEqual(self.window.packet_spin.value(), 56)

    def test_stop_keeps_controls_disabled_until_worker_finishes(self):
        class PendingTask:
            def __init__(self):
                self.stop_requested = False

            def request_stop(self):
                self.stop_requested = True

        run_id = 17
        task = PendingTask()
        self.window._run_id = run_id
        self.window.is_monitoring = True
        self.window.probe_task = task

        self.window.stop_monitoring()

        self.assertTrue(task.stop_requested)
        self.assertFalse(self.window.host_entry.isEnabled())
        self.assertEqual(self.window.status_pill.text(), "STOPPING")

        self.window._handle_probe_task_finished(run_id)

        self.assertTrue(self.window.host_entry.isEnabled())
        self.assertEqual(self.window.status_pill.text(), "IDLE")

    def test_import_hosts_reads_all_targets_from_one_comma_separated_line(self):
        host_file = Path(self.temporary_directory.name) / "hosts.txt"
        host_file.write_text(
            "one.example, two.example; Router (192.0.2.10), 192.0.2.11\n",
            encoding="utf-8",
        )
        self.window.host_entry.clear()
        self.window._open_file_dialog = lambda *args, **kwargs: str(host_file)

        self.window.import_hosts_from_file()

        self.assertEqual(
            self.window.host_entry.text(),
            "one.example, two.example, 192.0.2.10, 192.0.2.11",
        )

    def test_supplied_500_host_list_imports_without_changing_targets(self):
        host_file = PROJECT_DIR / "ping_domains.txt"
        self.window.host_entry.clear()
        self.window._open_file_dialog = lambda *args, **kwargs: str(host_file)

        self.window.import_hosts_from_file()

        imported = self.window.host_entry.text().split(", ")
        self.assertEqual(len(imported), 500)
        self.assertEqual(imported[0], "google.com")
        self.assertIn("hotels.com", imported)

    def test_result_updates_are_main_thread_owned_and_history_is_bounded(self):
        target = "localhost"
        self.window._run_id = 4
        self.window.is_monitoring = True
        self.window.graph_history_points = 2
        self.window.host_data[target] = {
            "latencies": deque(maxlen=2),
            "latency_count": 0,
            "latency_sum": 0.0,
            "min_latency": None,
            "max_latency": None,
            "sent": 0,
            "received": 0,
            "timeline_stats": [],
            "log_rows": [],
            "last_status": None,
            "last_error": None,
            "consecutive_fails": 0,
            "last_down_time": "-",
        }
        self.window.table.setRowCount(1)
        host_item = QTableWidgetItem(target)
        self.window.table.setItem(0, 0, host_item)
        self.window._host_items[target] = host_item

        for latency in (10.0, 20.0, 30.0):
            self.window._handle_probe_result(4, target, ProbeResult("Connected", latency))

        data = self.window.host_data[target]
        self.assertEqual(list(data["latencies"]), [20.0, 30.0])
        self.assertEqual(ping_monitor.PingMonitorApp._compute_host_stats(data)[3:], (10.0, 30.0, 20.0))

        self.window._handle_probe_result(3, target, ProbeResult("Connected", 40.0))
        self.assertEqual(data["sent"], 3)

    @patch("probes.async_ping", new_callable=AsyncMock)
    def test_thread_pool_probe_updates_the_window(self, ping_mock):
        ping_mock.return_value = SimpleNamespace(is_alive=True, avg_rtt=0.5)
        self.window.host_entry.setText("127.0.0.1")
        self.window.timeout_spin.setValue(0.5)
        self.window.ping_interval_seconds = 1.0
        self.window.toggle_monitoring()

        event_loop = QEventLoop()
        QTimer.singleShot(500, event_loop.quit)
        event_loop.exec()
        self.window.stop_monitoring()

        data = self.window.host_data["127.0.0.1"]
        self.assertEqual(data["sent"], 1)
        self.assertEqual(data["received"], 1)
        self.assertEqual(self.window.table.item(0, 1).text(), "Connected")
        ping_mock.assert_awaited_once()

    def test_csv_export_includes_summary_down_and_probe_rows(self):
        target = "invalid.example"
        self.window.host_data[target] = {
            "sent": 1,
            "received": 0,
            "latencies": deque(maxlen=60),
            "latency_count": 0,
            "latency_sum": 0.0,
            "min_latency": None,
            "max_latency": None,
            "last_status": "No ICMP Reply",
            "last_error": "name lookup failed",
            "consecutive_fails": 1,
            "last_down_time": "2026-10-03 12:00:00",
            "log_rows": [(1, "2026-10-03 12:00:00", "No ICMP Reply", "")],
        }
        report_path = Path(self.temporary_directory.name) / "report.csv"
        self.window._open_file_dialog = lambda *args, **kwargs: str(report_path)

        self.window.export_report_csv()

        with report_path.open(encoding="utf-8-sig", newline="") as report_file:
            rows = list(csv.DictReader(report_file))
        sections = [row["Section"] for row in rows if row["Section"]]
        self.assertEqual(sections, ["Summary", "No ICMP Reply", "Probe"])
        no_reply_row = next(row for row in rows if row["Section"] == "No ICMP Reply")
        self.assertEqual(no_reply_row["Status"], "No ICMP Reply")

    def test_csv_export_includes_recovered_hosts_with_lost_pings(self):
        target = "recovered.example"
        self.window.host_data[target] = {
            "sent": 3, "received": 2, "checks": 3,
            "latencies": deque([10.0, 0, 12.0], maxlen=60), "latency_count": 2,
            "latency_sum": 22.0, "min_latency": 10.0, "max_latency": 12.0,
            "last_status": "Connected", "last_error": None,
            "consecutive_fails": 0, "last_down_time": "2026-10-03 12:00:01",
            "timeline_stats": [
                "[2026-10-03 12:00:00] Seq 1: Success — Latency 10.00 ms",
                "[2026-10-03 12:00:01] Seq 2: FAILED — No ICMP Reply",
                "[2026-10-03 12:00:02] Seq 3: Success — Latency 12.00 ms",
            ],
            "log_rows": [
                (1, "2026-10-03 12:00:00", "Success", 10.0),
                (2, "2026-10-03 12:00:01", "No ICMP Reply", ""),
                (3, "2026-10-03 12:00:02", "Success", 12.0),
            ],
        }
        report_path = Path(self.temporary_directory.name) / "report_recovered.csv"
        self.window._open_file_dialog = lambda *args, **kwargs: str(report_path)

        self.window.export_report_csv()

        with report_path.open(encoding="utf-8-sig", newline="") as report_file:
            rows = list(csv.DictReader(report_file))

        # 1. Recovered host with lost pings must appear in No ICMP Reply section
        no_reply_rows = [r for r in rows if r["Section"] == "No ICMP Reply"]
        self.assertEqual(len(no_reply_rows), 1)
        self.assertEqual(no_reply_rows[0]["Host"], target)
        self.assertIn("Recovered", no_reply_rows[0]["Status"])
        self.assertEqual(no_reply_rows[0]["Last_Down_At"], "2026-10-03 12:00:01")

        # 2. Probe section must contain the No ICMP Reply row with diagnostic details
        probe_rows = [r for r in rows if r["Section"] == "Probe"]
        self.assertEqual(len(probe_rows), 3)
        lost_ping_probe = next(r for r in probe_rows if r["Status"] == "No ICMP Reply")
        self.assertEqual(lost_ping_probe["Sequence"], "2")
        self.assertEqual(lost_ping_probe["Details"], "No ICMP reply before timeout")

    def test_txt_export_includes_recovered_hosts_with_lost_pings(self):
        target = "recovered.example"
        self.window.host_data[target] = {
            "sent": 3, "received": 2, "checks": 3,
            "latencies": deque([10.0, 0, 12.0], maxlen=60), "latency_count": 2,
            "latency_sum": 22.0, "min_latency": 10.0, "max_latency": 12.0,
            "last_status": "Connected", "last_error": None,
            "consecutive_fails": 0, "last_down_time": "2026-10-03 12:00:01",
            "timeline_stats": [
                "[2026-10-03 12:00:00] Seq 1: Success — Latency 10.00 ms",
                "[2026-10-03 12:00:01] Seq 2: FAILED — No ICMP Reply",
                "[2026-10-03 12:00:02] Seq 3: Success — Latency 12.00 ms",
            ],
            "log_rows": [
                (1, "2026-10-03 12:00:00", "Success", 10.0),
                (2, "2026-10-03 12:00:01", "No ICMP Reply", ""),
                (3, "2026-10-03 12:00:02", "Success", 12.0),
            ],
        }
        report_path = Path(self.temporary_directory.name) / "report_recovered.txt"
        self.window._open_file_dialog = lambda *args, **kwargs: str(report_path)

        self.window.export_report_summary()

        report = report_path.read_text(encoding="utf-8")
        self.assertNotIn("No target recorded an ICMP timeout", report)
        self.assertIn("recovered.example", report)
        self.assertIn("Recovered", report)
        self.assertIn("2026-10-03 12:00:01", report)

    def test_xlsx_export_probe_details_includes_no_icmp_reply_description(self):
        target = "recovered.example"
        self.window.host_data[target] = {
            "sent": 2, "received": 1, "checks": 2,
            "latencies": deque([10.0, 0], maxlen=60), "latency_count": 1,
            "latency_sum": 10.0, "min_latency": 10.0, "max_latency": 10.0,
            "last_status": "Connected", "last_error": None,
            "consecutive_fails": 0, "last_down_time": "2026-10-03 12:00:01",
            "timeline_stats": [],
            "log_rows": [
                (1, "2026-10-03 12:00:00", "Success", 10.0),
                (2, "2026-10-03 12:00:01", "No ICMP Reply", ""),
            ],
        }
        report_path = Path(self.temporary_directory.name) / "report_recovered.xlsx"
        self.window._open_file_dialog = lambda *args, **kwargs: str(report_path)

        self.window.export_report_xlsx()

        workbook = load_workbook(report_path, data_only=True)
        probe_sheet = workbook["Probe Details"]
        rows = list(probe_sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Host", "Sequence", "Timestamp", "Status", "Latency ms", "Details"))
        lost_ping_row = next(r for r in rows[1:] if r[3] == "No ICMP Reply")
        self.assertEqual(lost_ping_row[0], target)
        self.assertEqual(lost_ping_row[1], 2)
        self.assertEqual(lost_ping_row[5], "No ICMP reply before timeout")

    def test_in_flight_probes_processed_during_stop_pending(self):
        target = "pending.example"
        self.window._run_id = 55
        self.window.is_monitoring = False
        self.window._stop_pending = True
        self.window.host_data[target] = {
            "sent": 0, "received": 0, "checks": 0,
            "latencies": deque(maxlen=60), "latency_count": 0,
            "latency_sum": 0.0, "min_latency": None, "max_latency": None,
            "timeline_stats": [], "log_rows": [], "last_status": None,
            "last_error": None, "consecutive_fails": 0, "last_down_time": "-",
        }
        self.window.table.setRowCount(1)
        item = QTableWidgetItem(target)
        self.window.table.setItem(0, 0, item)
        self.window._host_items[target] = item

        batch = [(target, ProbeResult("No ICMP Reply"))]
        self.window._handle_probe_batch(55, batch)

        self.assertEqual(len(self.window.host_data[target]["log_rows"]), 1)
        self.assertEqual(self.window.host_data[target]["log_rows"][0][2], "No ICMP Reply")
        self.assertEqual(self.window.host_data[target]["last_status"], "No ICMP Reply")
        self.assertNotEqual(self.window.host_data[target]["last_down_time"], "-")

    def test_xlsx_export_contains_summary_and_detail_sheets(self):
        target = "router.example"
        self.window.host_data[target] = {
            "sent": 2, "received": 1, "checks": 3,
            "latencies": deque([5.0, 9.0], maxlen=60), "latency_count": 2,
            "latency_sum": 14.0, "min_latency": 5.0, "max_latency": 9.0,
            "last_status": "No ICMP Reply", "last_error": None,
            "consecutive_fails": 1, "last_down_time": "2026-10-03 12:00:00",
            "log_rows": [(1, "12:00:00", "Success", 5.0),
                         (2, "12:00:01", "Success", 9.0),
                         (3, "12:00:02", "No ICMP Reply", "")],
        }
        report_path = Path(self.temporary_directory.name) / "report.xlsx"
        self.window._open_file_dialog = lambda *args, **kwargs: str(report_path)
        messages = []
        self.window._msg = lambda *args: messages.append(args)

        self.window.export_report_xlsx()

        self.assertTrue(report_path.is_file(), messages)
        workbook = load_workbook(report_path, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Summary Info", "Host Details", "Probe Details"])
        summary = {row[0]: row[1] for row in workbook["Summary Info"].iter_rows(min_row=5, values_only=True)}
        self.assertEqual(summary["Total targets"], 1)
        self.assertEqual(summary["Targets with no ICMP reply"], 1)
        self.assertEqual(summary["ICMP requests sent"], 2)
        self.assertEqual(summary["ICMP replies received"], 1)
        self.assertEqual(summary["ICMP packet loss (%)"], 50.0)
        self.assertEqual(workbook["Host Details"].max_row, 2)
        self.assertEqual(workbook["Probe Details"].max_row, 4)

    def test_txt_export_starts_with_summary_overview(self):
        target = "router.example"
        self.window.host_data[target] = {
            "sent": 2, "received": 1, "checks": 3,
            "latencies": deque([5.0, 9.0], maxlen=60), "latency_count": 2,
            "latency_sum": 14.0, "min_latency": 5.0, "max_latency": 9.0,
            "last_status": "No ICMP Reply", "last_error": None,
            "consecutive_fails": 1, "last_down_time": "2026-10-03 12:00:00",
            "timeline_stats": ["Seq 1: Connected"], "log_rows": [],
        }
        report_path = Path(self.temporary_directory.name) / "report.txt"
        self.window._open_file_dialog = lambda *args, **kwargs: str(report_path)

        self.window.export_report_summary()

        report = report_path.read_text(encoding="utf-8")
        self.assertLess(report.index("[1. SUMMARY INFO]"), report.index("[2. HOST AGGREGATE METRICS]"))
        self.assertIn("Total targets", report)
        self.assertIn("Targets with no ICMP reply", report)
        self.assertIn("ICMP packet loss (%)", report)

    def test_export_default_filenames_include_timestamp(self):
        target = "router.example"
        self.window.host_data[target] = {
            "sent": 1, "received": 1, "checks": 1,
            "latencies": deque([5.0], maxlen=60), "latency_count": 1,
            "latency_sum": 5.0, "min_latency": 5.0, "max_latency": 5.0,
            "last_status": "Connected", "last_error": None,
            "consecutive_fails": 0, "last_down_time": "-",
            "timeline_stats": ["Seq 1: Connected"], "log_rows": [(1, "12:00:00", "Success", 5.0)],
        }
        captured_defaults = {}

        def fake_dialog(title, name_filter, save=False, default_name="", default_suffix=""):
            captured_defaults[default_suffix] = default_name
            return None

        self.window._open_file_dialog = fake_dialog

        self.window.export_report_summary()
        self.window.export_report_xlsx()
        self.window.export_report_csv()

        self.assertTrue(re.match(r"^network_report_\d{8}_\d{6}\.txt$", captured_defaults.get("txt", "")))
        self.assertTrue(re.match(r"^network_report_\d{8}_\d{6}\.xlsx$", captured_defaults.get("xlsx", "")))
        self.assertTrue(re.match(r"^network_report_\d{8}_\d{6}\.csv$", captured_defaults.get("csv", "")))

    def test_dns_errors_do_not_count_as_sent_packets_or_host_downs(self):
        target = "nonexistent.example"
        self.window._run_id = 3
        self.window.is_monitoring = True
        self.window.host_data[target] = {
            "sent": 0, "received": 0, "checks": 0,
            "latencies": deque(maxlen=60), "latency_count": 0,
            "latency_sum": 0.0, "min_latency": None, "max_latency": None,
            "timeline_stats": [], "log_rows": [], "last_status": None,
            "last_error": None, "consecutive_fails": 0, "last_down_time": "-",
        }
        self.window.table.setRowCount(1)
        host_item = QTableWidgetItem(target)
        self.window.table.setItem(0, 0, host_item)
        self.window._host_items[target] = host_item

        self.window._handle_probe_result(
            3, target, ProbeResult("DNS Error", error="lookup failed", packet_sent=False)
        )

        self.assertEqual(self.window.host_data[target]["sent"], 0)
        self.assertEqual(self.window.host_data[target]["checks"], 1)
        self.assertFalse(self.window._is_down_status("DNS Error"))
        self.window.refresh_dashboard()
        self.assertIn("DNS/probe issue", self.window.down_summary_label.text())
        self.assertNotIn("awaiting first result", self.window.down_summary_label.text())
        self.assertEqual(self.window.host_count_label.text(), "0 up / 0 monitored")
        self.assertEqual(self.window.table.rowCount(), 0)

    def test_no_reply_count_matches_sorted_table_and_other_issue_counts(self):
        def host_data(status, received=0):
            return {
                "sent": 1, "received": received, "checks": 1,
                "latencies": deque(maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": status,
                "last_error": None, "consecutive_fails": 1 if status != "Connected" else 0,
                "last_down_time": "2026-10-03 12:00:00" if status == "No ICMP Reply" else "-",
            }

        self.window.host_data = {
            "z-down.example": host_data("No ICMP Reply"),
            "a-down.example": host_data("No ICMP Reply"),
            "b-dns.example": host_data("DNS Error"),
            "c-up.example": host_data("Connected", received=1),
        }
        self.window.down_table.sortItems(0, Qt.DescendingOrder)

        self.window.refresh_dashboard()

        self.assertEqual(self.window.down_table.rowCount(), 2)
        self.assertIn("2 of 4 down", self.window.down_summary_label.text())
        self.assertIn("1 DNS/probe issue", self.window.down_summary_label.text())
        self.assertEqual(self.window.host_count_label.text(), "1 up / 3 monitored")
        displayed = {
            self.window.down_table.item(row, 0).text(): self.window.down_table.item(row, 1).text()
            for row in range(self.window.down_table.rowCount())
        }
        self.assertEqual(set(displayed), {"z-down.example", "a-down.example"})
        self.assertTrue(all(status == "No ICMP Reply" for status in displayed.values()))

        self.window.host_data["z-down.example"]["last_status"] = "Connected"
        self.window.refresh_dashboard()

        self.assertEqual(self.window.down_table.rowCount(), 1)
        self.assertIn("1 of 4 down", self.window.down_summary_label.text())
        self.assertEqual(self.window.host_count_label.text(), "2 up / 3 monitored")

    def test_handle_probe_batch_updates_table_and_host_data(self):
        self.window._run_id = 99
        self.window.is_monitoring = True
        self.window.host_data = {
            "node1.example": {
                "sent": 0, "received": 0, "checks": 0,
                "latencies": deque(maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": None,
                "last_error": None, "consecutive_fails": 0,
                "last_down_time": "-",
            },
            "node2.example": {
                "sent": 0, "received": 0, "checks": 0,
                "latencies": deque(maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": None,
                "last_error": None, "consecutive_fails": 0,
                "last_down_time": "-",
            },
        }
        self.window.table.setRowCount(2)
        for r, host in enumerate(["node1.example", "node2.example"]):
            item = QTableWidgetItem(host)
            self.window.table.setItem(r, 0, item)
            self.window._host_items[host] = item

        batch = [
            ("node1.example", ProbeResult("Connected", 15.5)),
            ("node2.example", ProbeResult("No ICMP Reply")),
        ]
        self.window._handle_probe_batch(99, batch)

        self.assertEqual(self.window.host_data["node1.example"]["received"], 1)
        self.assertEqual(self.window.host_data["node2.example"]["consecutive_fails"], 1)
        self.assertEqual(self.window.table.item(0, 1).text(), "Connected")
        self.assertEqual(self.window.table.item(1, 1).text(), "No ICMP Reply")

    def test_handle_probe_batch_removes_dns_error_host_from_table(self):
        self.window._run_id = 101
        self.window.is_monitoring = True
        self.window.host_data = {
            "node1.example": {
                "sent": 0, "received": 0, "checks": 0,
                "latencies": deque(maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": None,
                "last_error": None, "consecutive_fails": 0,
                "last_down_time": "-",
            },
            "node2_bad.example": {
                "sent": 0, "received": 0, "checks": 0,
                "latencies": deque(maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": None,
                "last_error": None, "consecutive_fails": 0,
                "last_down_time": "-",
            },
            "node3.example": {
                "sent": 0, "received": 0, "checks": 0,
                "latencies": deque(maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": None,
                "last_error": None, "consecutive_fails": 0,
                "last_down_time": "-",
            },
        }
        self.window.table.setRowCount(3)
        for r, host in enumerate(["node1.example", "node2_bad.example", "node3.example"]):
            item = QTableWidgetItem(host)
            self.window.table.setItem(r, 0, item)
            self.window._host_items[host] = item

        batch = [
            ("node1.example", ProbeResult("Connected", 12.0)),
            ("node2_bad.example", ProbeResult("DNS Error", error="Domain unresolved", packet_sent=False)),
            ("node3.example", ProbeResult("Connected", 18.0)),
        ]
        self.window._handle_probe_batch(101, batch)

        # The DNS Error host must be removed from the table immediately
        self.assertEqual(self.window.table.rowCount(), 2)
        remaining = [self.window.table.item(r, 0).text() for r in range(self.window.table.rowCount())]
        self.assertIn("node1.example", remaining)
        self.assertIn("node3.example", remaining)
        self.assertNotIn("node2_bad.example", remaining)
        self.assertNotIn("node2_bad.example", self.window._host_items)

        # But data must be retained in host_data for the DNS inspector
        self.assertEqual(self.window.host_data["node2_bad.example"]["last_status"], "DNS Error")

    def test_chart_excludes_dns_error_hosts_and_updates_dns_pill(self):
        self.window.host_data = {
            "good1.example": {
                "sent": 1, "received": 1, "checks": 1,
                "latencies": deque([12.5], maxlen=60), "latency_count": 1,
                "latency_sum": 12.5, "min_latency": 12.5, "max_latency": 12.5,
                "timeline_stats": [], "log_rows": [], "last_status": "Connected",
                "last_error": None, "consecutive_fails": 0, "last_down_time": "-",
            },
            "good2.example": {
                "sent": 1, "received": 0, "checks": 1,
                "latencies": deque([0], maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": "No ICMP Reply",
                "last_error": None, "consecutive_fails": 1, "last_down_time": "-",
            },
            "dns_bad1.example": {
                "sent": 0, "received": 0, "checks": 1,
                "latencies": deque([], maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": "DNS Error",
                "last_error": "Domain not found", "consecutive_fails": 1, "last_down_time": "-",
            },
            "dns_bad2.example": {
                "sent": 0, "received": 0, "checks": 1,
                "latencies": deque([], maxlen=60), "latency_count": 0,
                "latency_sum": 0.0, "min_latency": None, "max_latency": None,
                "timeline_stats": [], "log_rows": [], "last_status": "DNS Error",
                "last_error": "Name or service not known", "consecutive_fails": 1, "last_down_time": "-",
            },
        }
        self.window.table.setRowCount(4)
        for r, host in enumerate(["good1.example", "good2.example", "dns_bad1.example", "dns_bad2.example"]):
            item = QTableWidgetItem(host)
            self.window.table.setItem(r, 0, item)
            self.window._host_items[host] = item

        self.window.refresh_dashboard()

        # Telemetry chart must only contain valid pingable (Connected) and timeout (No ICMP Reply) hosts
        self.assertIn("good1.example", self.window.chart.history)
        self.assertIn("good2.example", self.window.chart.history)
        self.assertNotIn("dns_bad1.example", self.window.chart.history)
        self.assertNotIn("dns_bad2.example", self.window.chart.history)

        # DNS Error hosts must be removed from the Live Telemetry table
        self.assertEqual(self.window.table.rowCount(), 2)
        table_hosts = [
            self.window.table.item(r, 0).text() for r in range(self.window.table.rowCount())
        ]
        self.assertIn("good1.example", table_hosts)
        self.assertIn("good2.example", table_hosts)
        self.assertNotIn("dns_bad1.example", table_hosts)
        self.assertNotIn("dns_bad2.example", table_hosts)
        self.assertEqual(self.window.host_count_label.text(), "1 up / 2 monitored")

        # Toolbar DNS pill badge must reflect the 2 DNS failures
        self.assertEqual(self.window.dns_issues_btn.text(), "⚠ 2 DNS Issues")

        # Clearing history resets the pill
        self.window._confirm = lambda *args: True
        self.window.clear_history()
        self.assertEqual(self.window.dns_issues_btn.text(), "0 DNS Issues")

    def test_dns_issues_dialog_table_filtering_and_copy(self):
        issues = [
            ("unresolved1.example", {"last_error": "No address found", "last_down_time": "2026-10-18 10:00:00"}),
            ("unresolved2.example", {"last_error": "Server failure", "last_down_time": "2026-10-18 10:01:00"}),
        ]
        dialog = ping_monitor.DNSIssuesDialog(issues, dark=True)
        self.assertEqual(dialog.table.rowCount(), 2)
        self.assertEqual(dialog.table.item(0, 0).text(), "unresolved1.example")
        self.assertEqual(dialog.table.item(1, 0).text(), "unresolved2.example")

        # Test search filter
        dialog._filter_table("unresolved1")
        self.assertFalse(dialog.table.isRowHidden(0))
        self.assertTrue(dialog.table.isRowHidden(1))
        self.assertIn("Showing 1 of 2", dialog.counter_label.text())

        dialog._filter_table("")
        self.assertFalse(dialog.table.isRowHidden(0))
        self.assertFalse(dialog.table.isRowHidden(1))
        self.assertEqual(dialog.counter_label.text(), "Total: 2 hosts")

        # Test clipboard copy
        dialog.copy_hostnames()
        clipboard_text = QApplication.clipboard().text()
        self.assertIn("unresolved1.example", clipboard_text)
        self.assertIn("unresolved2.example", clipboard_text)
        self.assertIn("Copied 2 hostnames", dialog.feedback_label.text())

    def test_packet_loss_sorting_and_row_integrity(self):
        hosts = [f"node_{i}.net" for i in range(1, 21)]
        self.window.hosts = hosts
        self.window.host_entry.setText(", ".join(hosts))
        self.window.is_monitoring = True
        self.window._run_id = 42

        for i, host in enumerate(hosts):
            self.window.host_data[host] = {
                "sent": 10, "received": 10 - (i % 5) * 2, "checks": 10,
                "latencies": deque([15.0 + i], maxlen=60), "latency_count": 1, "latency_sum": 15.0 + i,
                "min_latency": 15.0 + i, "max_latency": 15.0 + i,
                "consecutive_fails": (i % 5), "last_status": "Connected",
                "last_error": None, "last_down_time": "-",
                "timeline_stats": [], "log_rows": [],
            }
            row = self.window.table.rowCount()
            self.window.table.insertRow(row)
            loss_pct = ((10 - self.window.host_data[host]["received"]) / 10) * 100
            vals = [host, "Connected", "10", f"{15.0+i:.1f} ms", f"{loss_pct:.1f}%", str(i % 5)]
            sort_vals = {2: 10, 3: 15.0 + i, 4: loss_pct, 5: i % 5}
            for col, val in enumerate(vals):
                it = ping_monitor.SortableTableWidgetItem(val, numeric=col in (2, 3, 4, 5), sort_value=sort_vals.get(col))
                self.window.table.setItem(row, col, it)
                if col == 0:
                    self.window._host_items[host] = it

        # Click Packet Loss header (col 4) - should default to Descending (highest loss first)
        self.window.on_table_header_clicked(4)
        self.application.processEvents()

        # Check top row has highest loss
        top_loss_item = self.window.table.item(0, 4)
        top_loss_val = top_loss_item.data(Qt.UserRole)
        self.assertEqual(top_loss_val, 80.0)

        # Send an update for node_1.net while sort is active
        # node_1 currently has loss 0%. Update it to No ICMP Reply
        self.window._handle_probe_result(42, "node_1.net", ProbeResult("No ICMP Reply", None))
        self.application.processEvents()

        # Row integrity: col 0 must match host_items lookup
        for host in hosts:
            it = self.window._host_items[host]
            row_idx = it.row()
            cell_host = self.window.table.item(row_idx, 0).text()
            self.assertEqual(cell_host, host)

        # Trigger dashboard refresh (which does periodic sort)
        self.window.refresh_dashboard()
        self.application.processEvents()

        # node_1.net now has loss, ensure it remains matched with its row
        it_1 = self.window._host_items["node_1.net"]
        row_1 = it_1.row()
        self.assertEqual(self.window.table.item(row_1, 0).text(), "node_1.net")

    def test_csv_sanitization_dde_attacks(self):
        # Formula injection strings must be prefixed with single quote
        self.assertEqual(ping_monitor._sanitize_csv_cell("=cmd|'/C calc'!A0"), "'=cmd|'/C calc'!A0")
        self.assertEqual(ping_monitor._sanitize_csv_cell("+cmd|'/C calc'!A0"), "'+cmd|'/C calc'!A0")
        self.assertEqual(ping_monitor._sanitize_csv_cell("-cmd|'/C calc'!A0"), "'-cmd|'/C calc'!A0")
        self.assertEqual(ping_monitor._sanitize_csv_cell("@SUM(A1:A10)"), "'@SUM(A1:A10)")
        self.assertEqual(ping_monitor._sanitize_csv_cell("\tmalicious"), "'\tmalicious")
        # Regular numeric values must remain unchanged
        self.assertEqual(ping_monitor._sanitize_csv_cell(12.5), 12.5)
        self.assertEqual(ping_monitor._sanitize_csv_cell("-5.2"), "-5.2")
        self.assertEqual(ping_monitor._sanitize_csv_cell("+10"), "+10")
        self.assertEqual(ping_monitor._sanitize_csv_cell("8.8.8.8"), "8.8.8.8")


if __name__ == "__main__":
    unittest.main()



